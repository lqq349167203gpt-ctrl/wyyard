from datetime import datetime, timezone
from io import BytesIO
from types import SimpleNamespace as NS

import pytest
from openpyxl import load_workbook

from app.api import customers as api
from app.services import customer_access_service, position_edit_permission_service


def customer(customer_id, **values):
    return NS(**{
        "id": customer_id, "nickname": customer_id, "name": "姓名", "gender": "女",
        "member_type": "体验会员", "card_status": "即将用完", "card_status_tone": "warning",
        "service_teacher": "老师甲", "referrer": "甲、乙", "referrers": ["甲", "乙"],
        "referrer_handler": "承接甲", "referral_date": "2026-10-01",
        "created_at": datetime(2026, 10, 1, tzinfo=timezone.utc), "created_by": "录入人",
        **values,
    })


@pytest.fixture
def export_data(monkeypatch):
    from app.api import customer_detail
    records = [customer(f"客户{i:02d}") for i in range(25)]
    logs = []
    monkeypatch.setattr(api.customer_service, "list_customers", lambda: records)
    monkeypatch.setattr(api, "_build_transaction_counts", lambda: {item.id: 2 for item in records})
    monkeypatch.setattr(api, "_count_customer_activities", lambda _: 0)
    monkeypatch.setattr(api, "customer_visit_summary", lambda ids: {key: (0, "") for key in ids})
    monkeypatch.setattr(api.customer_tag_service, "visible_tags_by_customer", lambda _: {
        item.id: [{"name": "可见客户标签"}] for item in records
    })
    monkeypatch.setattr(api.operation_log_service, "create_log", lambda data, **kwargs: logs.append((data, kwargs)))
    monkeypatch.setattr(customer_detail, "get_customer_detail", lambda *args, **kwargs: {
        "customer": {"total_payment": 0, "first_visit": ""},
    })
    return records, logs


def sheet_rows(response):
    assert response.status_code == 200, response.text
    sheet = load_workbook(BytesIO(response.content)).active
    headers = [cell.value for cell in sheet[1]]
    return sheet, [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]


def test_export_contains_all_pages_and_separate_labels_and_audit(client, export_data):
    response = client.get("/api/customers/export")
    sheet, rows = sheet_rows(response)
    assert len(rows) == 25
    headers = [cell.value for cell in sheet[1]]
    assert headers[headers.index("会员身份") + 1] == "卡次状态"
    assert rows[0]["卡次状态"] == "即将用完"
    assert rows[0]["客户标签"] == "可见客户标签"
    assert rows[0]["引流人"] == "甲、乙"
    assert rows[0]["到店次数"] == 0 and rows[0]["活动次数"] == 0
    assert rows[0]["交易笔数"] == 2
    options = client.get("/api/customers/export-options").json()
    assert headers == [column["label"] for column in options if column["default"]]
    assert "帖子名称" in headers and "部门身份" not in headers and "跟进阶段" not in headers
    assert not {"需求标签", "跟进节点", "跟进行动", "跟进计划", "创建时间", "更新时间", "累计消费"}.intersection(headers)
    assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref == sheet.dimensions
    assert "filename*=UTF-8''" in response.headers["Content-Disposition"]
    _, logs = export_data
    assert len(logs) == 1
    assert logs[0][1]["extra"]["method"] == "EXPORT"
    assert logs[0][1]["extra"]["after_data"]["客户数"] == 25


def test_export_defaults_select_all_permitted_columns_in_four_groups(client):
    options = client.get("/api/customers/export-options").json()
    default_groups = {"基础资料", "联系来源", "客户档案", "数据统计"}
    assert {column["group"] for column in options if column["default"]} == default_groups
    assert all(column["default"] == (column["group"] in default_groups) for column in options)
    assert not {"positions", "follow_up_status", "need_tags", "follow_up_node", "follow_up_action", "tracking_plan", "created_at", "updated_at", "total_payment"}.intersection(
        column["key"] for column in options
    )


def test_export_post_name_is_separate_from_content_link(client, export_data):
    export_data[0][:] = [customer("帖子客户", traffic_source="小红书", traffic_post_name="夏日放松分享",
                                traffic_source_detail="https://example.com/post/1")]
    sheet, rows = sheet_rows(client.get("/api/customers/export", params=[
        ("columns", "traffic_post_name"), ("columns", "traffic_source_detail"),
    ]))
    assert [cell.value for cell in sheet[1]] == ["帖子名称", "来源详情"]
    assert rows == [{"帖子名称": "夏日放松分享", "来源详情": "https://example.com/post/1"}]


@pytest.mark.parametrize("source", ["小红书", "抖音", "公众号", "视频号"])
def test_social_post_name_persists_and_returns_in_customer_views(client, created_customer, source):
    from app.services.storage import load_item
    cid = created_customer["id"]
    updated = client.patch(f"/api/customers/{cid}", json={
        "traffic_source": source, "traffic_post_name": f"{source}帖子标题",
        "traffic_source_detail": "https://example.com/post/123",
    })
    assert updated.status_code == 200, updated.text
    assert load_item("customers.json", cid)["traffic_post_name"] == f"{source}帖子标题"
    assert client.get(f"/api/customers/{cid}").json()["traffic_post_name"] == f"{source}帖子标题"
    rows = client.get("/api/customers").json()
    assert next(row for row in rows if row["id"] == cid)["traffic_post_name"] == f"{source}帖子标题"


def test_export_combines_six_filters_and_includes_both_date_boundaries(client, export_data):
    records, _ = export_data
    records[:] = [customer("开始当天"), customer("结束当天", referral_date="2026-10-20", member_type="30次卡")]
    for key, value in (("card_status", "卡次充足"), ("member_type", "未知身份"),
                       ("service_teacher", "老师甲乙"), ("referrers", ["乙丙"]),
                       ("referrer_handler", "承接甲乙"), ("referral_date", "2026-09-30"),
                       ("referral_date", "2026-10-21"), ("referral_date", "")):
        records.append(customer(f"不匹配-{key}-{value}", **{key: value}))
    response = client.get("/api/customers/export", params=[
        ("card_statuses", "即将用完"), ("member_types", "体验会员"), ("member_types", "30次卡"),
        ("service_teachers", "老师甲"), ("referrers", "乙"), ("referrer_handlers", "承接甲"),
        ("date_from", "2026-10-01"), ("date_to", "2026-10-20"),
    ])
    _, rows = sheet_rows(response)
    assert {row["客户昵称"] for row in rows} == {"开始当天", "结束当天"}


@pytest.mark.parametrize("params", [
    {"date_from": "2026-10-20", "date_to": "2026-10-01"}, {"date_from": "日期错误"},
])
def test_export_rejects_invalid_date_ranges_without_logging(client, export_data, params):
    assert client.get("/api/customers/export", params=params).status_code == 422
    assert export_data[1] == []


def test_export_empty_result_returns_clear_message(client, export_data):
    response = client.get("/api/customers/export", params={"card_statuses": "不匹配"})
    assert response.status_code == 400
    assert "没有符合条件的客户" in response.json()["detail"]
    assert export_data[1] == []


def test_export_respects_related_scope_and_transaction_permissions(client, export_data, monkeypatch):
    records, _ = export_data
    records[:] = [customer("可见", referrers=["甲"]), customer("不可见", referrers=["外人"])]
    roles = ["普通员工"]
    monkeypatch.setattr(api, "get_request_roles", lambda _: roles)
    monkeypatch.setattr(customer_access_service, "get_request_roles", lambda _: roles)
    permissions = position_edit_permission_service._empty_customer_access()
    permissions.update(scope="related", transaction_access="none")
    permissions["relations"]["referrer"] = True
    monkeypatch.setattr(customer_access_service, "get_customer_permissions", lambda _: permissions)
    monkeypatch.setattr(customer_access_service, "actor_name", lambda _: "甲")
    sheet, rows = sheet_rows(client.get("/api/customers/export"))
    assert [row["客户昵称"] for row in rows] == ["可见"]
    assert "交易笔数" not in [cell.value for cell in sheet[1]]


def test_export_neutralizes_customer_formula_input(client, export_data):
    export_data[0][:] = [customer("=1+1", name="+SUM(1,2)", referrers=["@人"])]
    sheet, rows = sheet_rows(client.get("/api/customers/export"))
    assert rows[0]["客户昵称"] == "'=1+1"
    assert rows[0]["姓名"] == "'+SUM(1,2)"
    assert rows[0]["引流人"] == "'@人"
    assert all(cell.data_type != "f" for row in sheet.iter_rows() for cell in row)


def test_export_only_checked_columns_in_requested_order(client, export_data):
    response = client.get("/api/customers/export", params=[("columns", "referral_date"), ("columns", "card_status")])
    sheet, rows = sheet_rows(response)
    assert [cell.value for cell in sheet[1]] == ["引流日期", "卡次状态"]
    assert rows[0] == {"引流日期": "2026-10-01", "卡次状态": "即将用完"}
    assert client.get("/api/customers/export", params={"columns": ""}).status_code == 422


def test_available_columns_and_export_respect_hidden_profile_and_payment_fields(client, export_data, monkeypatch):
    access = position_edit_permission_service._empty_customer_access()
    access["scope"] = "all"
    monkeypatch.setattr(customer_access_service, "get_customer_permissions", lambda _: access)
    options = client.get("/api/customers/export-options")
    assert options.status_code == 200
    keys = {column["key"] for column in options.json()}
    assert {"nickname", "card_status", "referral_date"} <= keys
    assert not {"basic_info", "assessment", "work_description", "payment_records", "total_payment", "follow_up", "tracking_plan"} & keys
    assert client.get("/api/customers/export", params={"columns": "basic_info"}).status_code == 403


def test_contact_export_obeys_copy_permission(client, export_data, monkeypatch):
    from app.services import customer_contact_service
    export_data[0][:] = [customer("带联系方式", phone="13812345678", wechat="wechat888")]
    monkeypatch.setattr(customer_contact_service, "can_access", lambda *args: False)
    _, rows = sheet_rows(client.get("/api/customers/export", params=[("columns", "phone"), ("columns", "wechat")]))
    assert rows[0]["手机号（脱敏）"] == "138****5678"
    assert rows[0]["微信号（脱敏）"] != "wechat888"
    monkeypatch.setattr(customer_contact_service, "can_access", lambda *args: True)
    _, rows = sheet_rows(client.get("/api/customers/export", params={"columns": "phone"}))
    assert rows[0]["手机号"] == "13812345678"


def test_selected_follow_up_keeps_full_content_author_and_time_in_own_sheet(client, export_data, monkeypatch):
    from app.api import customer_detail
    export_data[0][:] = [customer("跟进客户")]
    monkeypatch.setattr(customer_detail, "get_customer_detail", lambda *args, **kwargs: {
        "visit_records": [{"visit_date": "2026-10-20", "visit_notes": [
            {"category": "follow_up", "content": "完整跟进内容", "feedback_person": "反馈人甲", "created_by": "记录人乙", "created_at": "2026-10-20 10:00"},
            {"category": "customer_info", "content": "未勾选的深度需求"},
        ]}], "activity_participant_notes": [],
    })
    response = client.get("/api/customers/export", params=[("columns", "nickname"), ("columns", "follow_up")])
    sheet_rows(response)
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == ["客户资料", "跟进记录"]
    rows = list(workbook["跟进记录"].values)
    assert "反馈人" in rows[0] and "录入时间" in rows[0]
    assert "完整跟进内容" in rows[1] and "反馈人甲" in rows[1] and "记录人乙" in rows[1]
    assert "未勾选的深度需求" not in str(rows)


def test_card_export_keeps_purchase_details_and_respects_hidden_amounts(client, export_data, monkeypatch):
    from app.api import customer_detail
    export_data[0][:] = [customer("有卡客户")]
    monkeypatch.setattr(customer_detail, "get_customer_detail", lambda *args, **kwargs: {
        "purchase_summary": [{"type": "会员卡", "name": "30次卡", "remaining": 12, "total_amount": 398,
                              "debt_count": 2, "purchases": [{"deal_date": "2026-10-01", "purchase_count": 30, "amount": 398}]}],
    })
    response = client.get("/api/customers/export", params={"columns": "purchase_summary"})
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == ["客户资料"]
    assert "会员卡剩余次数" in next(workbook.active.values)
    access = position_edit_permission_service._empty_customer_access()
    access["scope"] = "all"
    access["detail_tabs"]["card_statistics"] = True
    monkeypatch.setattr(customer_access_service, "get_customer_permissions", lambda _: access)
    response = client.get("/api/customers/export", params={"columns": "purchase_summary"})
    workbook = load_workbook(BytesIO(response.content))
    assert "总金额" not in next(workbook.active.values)
    assert workbook.sheetnames == ["客户资料"]
    assert 12 in list(workbook.active.values)[1]
