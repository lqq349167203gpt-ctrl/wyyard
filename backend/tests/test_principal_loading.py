"""组织统计的排队、按需明细和旧端兼容回归。"""

import asyncio
import threading
from concurrent.futures import ThreadPoolExecutor
from copy import deepcopy
from types import SimpleNamespace as NS

import pytest
from pydantic import ValidationError

from app.middleware.course_deduction_consistency import CourseDeductionConsistencyMiddleware
from app.models.principal import PrincipalQuery
from app.services import principal_query_cache as cache
from app.services.principal_mobile_service import compact_breakdown, overview_metrics
from app.services.principal_response_service import detail_page


def test_principal_reads_bypass_busy_write_lock_but_course_writes_still_wait(monkeypatch):
    from app.services import project_deduction_service

    monkeypatch.setattr(project_deduction_service, "list_deductions", lambda: [])

    async def scenario():
        entered = []

        async def app(scope, receive, send):
            entered.append(scope["path"])

        middleware = CourseDeductionConsistencyMiddleware(app)
        await middleware.lock.acquire()
        write = asyncio.create_task(middleware({"type": "http", "method": "PATCH", "path": "/api/class-records/1"}, None, None))
        try:
            for path in ("/api/principal/query", "/api/principal/export"):
                await asyncio.wait_for(middleware({"type": "http", "method": "POST", "path": path}, None, None), timeout=0.2)
            assert entered == ["/api/principal/query", "/api/principal/export"]
            assert not write.done()
        finally:
            middleware.lock.release()
            await asyncio.wait_for(write, timeout=0.2)
        assert entered[-1] == "/api/class-records/1"

    asyncio.run(scenario())


def result_fixture():
    people = [{"id": f"c{i}", "name": f"客户{i}", "deals": i, "referrer": "老师甲",
               "arrive_count": 1, "identity": "体验会员", "tags": ["关注"],
               "visit_purpose": "长正文" * 100,
               "arrival_records": [{"id": f"v{i}", "arrive_date": "2026-09-01", "needs": "到店需求"},
                                   {"id": f"v{i}-2", "arrive_date": "2026-09-02", "needs": "第二次需求"}]}
              for i in range(25)]
    records = [{"id": f"v{i}", "customer_id": f"c{i % 5}", "date": "2026-09-01", "deals": i,
                "visit_purpose": "邀约隐私正文"} for i in range(25)]
    return {"items": [], "total": 0, "summary": {"引流人数": 25}, "breakdown": {
        "traffic": [{"key": "老师甲", "label": "老师甲", "count": 25, "customers": people}],
        "invite_inviters": [{"key": "老师乙", "label": "老师乙", "count": 5, "initiated_count": 25, "records": records}],
        "traffic_profile_fields": ["visit_purpose"],
    }}


def test_compact_summary_omits_unopened_details_and_preserves_unique_people():
    full = result_fixture()
    before = deepcopy(full)
    compact = compact_breakdown(full["breakdown"])
    assert compact["traffic"][0]["count"] == 25
    assert compact["traffic_profile_fields"] == ["visit_purpose"]
    assert "customers" not in compact["traffic"][0]
    assert "customer_ids" not in compact["invite_inviters"][0]
    metrics = overview_metrics(full["breakdown"], [])
    assert metrics["traffic_count"] == 25
    assert metrics["traffic_deals"] == sum(range(25))
    assert metrics["initiated_people"] == 5
    assert "records" not in compact["invite_inviters"][0]
    assert full == before


def test_traffic_page_sorts_all_rows_and_preserves_authorized_profile_fields():
    query = PrincipalQuery(overview_detail="traffic", detail_picks=["traffic:老师甲", "tag:关注"],
                           detail_sort_by="deals", detail_sort_order="desc", page=2, page_size=10)
    result = detail_page(result_fixture()["breakdown"], query)
    assert result["total"] == 25 and result["total_pages"] == 3
    assert [row["deals"] for row in result["items"]] == list(range(14, 4, -1))
    assert result["items"][0]["visit_purpose"] == "长正文" * 100
    assert "arrival_records" not in result["items"][0]
    query.detail_picks = ["traffic:不可见老师"]
    assert detail_page(result_fixture()["breakdown"], query)["total"] == 0


def test_arrival_pages_keep_each_visit_and_customer_modes_distinct():
    full = result_fixture()["breakdown"]
    query = PrincipalQuery(overview_detail="invite_arrivals", arrival_view="date", page=3, page_size=20)
    result = detail_page(full, query)
    assert result["total"] == 50 and len(result["items"]) == 10
    assert result["items"][0]["arrival_id"].startswith("v")
    assert result["items"][0]["id"].startswith("c")
    query.arrival_view = "customer"
    result = detail_page(full, query)
    assert result["total"] == 25 and result["page"] == 2


def test_inviter_details_only_read_requested_person_and_sort_before_page():
    query = PrincipalQuery(overview_detail="invite_initiated", overview_detail_key="老师乙",
                           detail_sort_by="deals", detail_sort_order="desc", page=2, page_size=10)
    result = detail_page(result_fixture()["breakdown"], query)
    assert result["total"] == 25
    assert [row["deals"] for row in result["items"]] == list(range(14, 4, -1))
    query.overview_detail_key = "其他人"
    assert detail_page(result_fixture()["breakdown"], query)["total"] == 0


@pytest.fixture
def cached_analysis(monkeypatch):
    cache._entries.clear()
    calls = []
    monkeypatch.setattr(cache, "data_revision", lambda: 0)
    monkeypatch.setattr(cache.position_edit_permission_service, "get_permissions", lambda _: {"scope": "all"})

    def analyze(*args, **kwargs):
        calls.append(True)
        return result_fixture()

    monkeypatch.setattr(cache.principal_service, "analyze", analyze)
    yield calls
    cache._entries.clear()


def request():
    return NS(state=NS(user_id="account", user_owner="老师甲", user_role="超级管理员"))


def test_summary_and_detail_pages_reuse_analysis_without_changing_legacy_or_mobile(cached_analysis):
    compact = cache.query_result(request(), PrincipalQuery(compact_overview=True))
    assert "records" not in compact["breakdown"]["invite_inviters"][0]
    detail = cache.query_result(request(), PrincipalQuery(overview_detail="traffic", page=2, page_size=20))
    assert len(detail["overview_detail"]["items"]) == 5
    assert set(detail) == {"overview_detail"}
    legacy = cache.query_result(request(), PrincipalQuery())
    assert len(legacy["breakdown"]["invite_inviters"][0]["records"]) == 25
    assert len(cached_analysis) == 1
    filtered = cache.query_result(request(), PrincipalQuery(compact_overview=True, detail_picks=["traffic:其他老师"]))
    assert filtered["overview_metrics"]["traffic_count"] == 0
    assert len(cached_analysis) == 1
    mobile = cache.query_result(request(), PrincipalQuery(compact_overview=True, mobile_group="traffic"))
    assert mobile["breakdown"]["traffic"][0]["customers"][0]["arrival_records"]


def test_concurrent_identical_queries_share_one_calculation(cached_analysis, monkeypatch):
    entered = threading.Event()
    release = threading.Event()

    def slow_analyze(*args, **kwargs):
        cached_analysis.append(True)
        entered.set()
        assert release.wait(timeout=2)
        return result_fixture()

    monkeypatch.setattr(cache.principal_service, "analyze", slow_analyze)
    with ThreadPoolExecutor(max_workers=2) as executor:
        first = executor.submit(cache.query_result, request(), PrincipalQuery(compact_overview=True))
        assert entered.wait(timeout=2)
        second = executor.submit(cache.query_result, request(), PrincipalQuery(overview_detail="traffic"))
        release.set()
        assert first.result(timeout=2)["summary"]["引流人数"] == 25
        assert second.result(timeout=2)["overview_detail"]["total"] == 25
    assert len(cached_analysis) == 1 and not cache._inflight


def test_invalid_detail_and_mobile_modes_are_rejected():
    with pytest.raises(ValidationError, match="按需明细"):
        PrincipalQuery(overview_detail="traffic", tab="conversion")
    with pytest.raises(ValidationError, match="按需明细"):
        PrincipalQuery(overview_detail="traffic", mobile_group="traffic")


def test_export_keeps_all_rows_and_matches_numeric_detail_sort(client, monkeypatch):
    import io

    from openpyxl import load_workbook

    monkeypatch.setattr(cache.principal_service, "analyze", lambda *_args, **_kwargs: result_fixture())
    response = client.post("/api/principal/export", json={
        "export_view": "traffic", "sort_by": "deals", "sort_order": "desc", "page_size": 10,
    })
    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert sheet.max_row == 26
    assert sheet["A2"].value == "客户24" and sheet["A26"].value == "客户0"


def test_shared_metrics_preserve_customer_filtering_rates_and_deduplication():
    full = result_fixture()["breakdown"]
    people = full["traffic"][0]["customers"]
    full["traffic_upsell_levels"] = [{"key": "trial", "label": "体验"}, {"key": "upgrade", "label": "升单"}]
    people[0].update(deals=1, upsell_levels=[{"key": "trial"}], products=[{"key": "membership", "label": "会员卡", "count": 1}],
                     cancel_count=1, no_show_count=2, arrive_count=3, invite_count=5)
    people[1].update(deals=2, is_upsell=True, upsell_levels=[{"key": "trial"}, {"key": "upgrade"}],
                     products=[{"key": "membership", "label": "会员卡", "count": 2}])
    full["invite_inviters"].append({"key": "老师丙", "initiated_count": 1, "records": [{"customer_id": "c0"}]})
    compact = compact_breakdown(full)
    assert compact["traffic"][0]["trial_count"] == 1
    assert compact["traffic"][0]["upsell_count"] == 1
    metrics = overview_metrics(full, ["traffic:老师甲", "tag:关注", "upsell:trial"])
    assert metrics["traffic_count"] == 1
    assert metrics["traffic_products"] == [{"key": "membership", "label": "会员卡", "count": 1}]
    assert metrics["invite"]["invited_count"] == {"times": 6, "people": 1}
    assert overview_metrics(full, [])["initiated_people"] == 5  # 多位邀约人邀约同一客户仍只记一人
    assert overview_metrics(full, ["inviter:老师丙"])["initiated_people"] == 1


@pytest.mark.parametrize("options", [
    {}, {"invite_view": "initiated"}, {"course_view": "participant"}, {"course_view": "teacher_follow_up"},
    {"tab": "orders", "include_overview": True}, {"tab": "orders", "list_view": "customer", "include_overview": True},
    {"tab": "courses"}, {"tab": "courses", "course_view": "participant"},
    {"tab": "courses", "course_view": "teacher_follow_up"}, {"tab": "conversion"},
])
def test_each_tab_compaction_keeps_rows_totals_and_summaries(client, options):
    legacy = client.post("/api/principal/query", json=options)
    compact = client.post("/api/principal/query", json={**options, "compact_overview": True})
    assert legacy.status_code == compact.status_code == 200
    before, after = legacy.json(), compact.json()
    for key in ("items", "total", "summary", "list_summary", "columns", "page", "total_pages"):
        assert after.get(key) == before.get(key), (options, key)
    assert after["overview_metrics"] == overview_metrics(before.get("breakdown") or {}, [])
    for group in after["breakdown"].get("traffic", []):
        assert "customers" not in group
    for group in after["breakdown"].get("invite_inviters", []):
        assert "records" not in group


def test_pc_export_applies_same_picks_and_quick_filter_as_detail(client, monkeypatch):
    import io

    from openpyxl import load_workbook

    full = result_fixture()
    monkeypatch.setattr(cache.principal_service, "analyze", lambda *_args, **_kwargs: full)
    query = PrincipalQuery(overview_detail="traffic", detail_picks=["traffic:老师甲"], detail_quick_filter="deals",
                           detail_sort_by="deals", detail_sort_order="desc", page_size=100)
    detail = detail_page(full["breakdown"], query)
    response = client.post("/api/principal/export", json={
        "export_view": "traffic", "breakdown": query.detail_picks, "detail_quick_filter": "deals",
        "sort_by": "deals", "sort_order": "desc", "page_size": 10,
    })
    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert [row[0] for row in sheet.iter_rows(min_row=2, values_only=True)] == [row["name"] for row in detail["items"]]
