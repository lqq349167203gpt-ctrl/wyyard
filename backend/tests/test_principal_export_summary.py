"""组织各卡片导出的顶部汇总与下拉统计共用服务端数据。"""

from copy import deepcopy
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.api import principal as api


@pytest.fixture
def export_data(monkeypatch):
    customer = {"id": "c1", "name": "客户甲", "referrer": "老师甲", "tags": [], "deals": 2,
                "invite_count": 3, "cancel_count": 1, "no_show_count": 1, "arrive_count": 1,
                "arrival_records": [{"arrive_date": "2026-10-01"}], "arrive_date": "2026-10-01"}
    result = {
        "summary": {"引流人数": 40, "交易笔数": 9, "成交人数": 5, "课程数": 6, "课时数": 12},
        "columns": [{"key": "date", "label": "课程日期"}, {"key": "name", "label": "课程"}],
        "items": [{"date": "2026-10-01", "name": "课程甲", "details": []},
                  {"date": "2026-10-02", "name": "课程乙", "details": []}],
        "total": 2,
        "breakdown": {
            "traffic": [{"key": "老师甲", "label": "老师甲", "count": 40, "customers": [customer]}],
            "traffic_filters": {key: [{"key": key, "label": label, "count": count}]
                                for key, label, count in (("identity", "年卡", 5), ("upsell", "已升单", 2),
                                                          ("stage", "未配置", 0), ("source", "自然到店", 3),
                                                          ("tag", "标签甲", 2))},
            "invite_traffic": [{"key": "老师甲", "label": "老师甲", "invite_count": 3, "customers": [customer]}],
            "invite_inviters": [{"key": "接待甲", "label": "接待甲", "initiated_count": 4,
                                 "records": [{"customer_id": "c1", "name": "客户甲", "tags": []}]}],
            "deals": [{"key": "membership", "label": "会员卡", "count": 6,
                       "subtypes": [{"key": "年卡", "label": "年卡", "count": 4},
                                    {"key": "体验卡", "label": "体验卡", "count": 2}]},
                      {"key": "other", "label": "其他项目", "count": 3, "subtypes": []}],
            "buys": [{"key": "首购", "label": "首购", "count": 5}, {"key": "复购", "label": "复购", "count": 4}],
            "courses": {
                "by_type": [{"key": "沙龙活动", "label": "沙龙活动", "count": 6, "hours": 12,
                             "subtypes": [{"key": "club｜读书会", "label": "俱乐部 · 读书会", "count": 6, "hours": 12}]}],
                # 多位老师可以共同上一堂课，不得把老师数量相加作为总课程数。
                "by_teacher": [{"key": "老师甲", "label": "老师甲", "count": 6, "hours": 12},
                               {"key": "老师乙", "label": "老师乙", "count": 6, "hours": 12}],
            },
            "participation": {
                "overall": {"count": 3, "people": 1, "owner_count": 1, "owner_people": 1,
                            "participant_count": 2, "participant_people": 1},
                "by_type": [{"key": "class", "label": "沙龙活动", "count": 2, "participant_count": 2,
                             "participant_people": 1, "subtypes": [{"label": "读书会", "count": 2}]}],
                "by_participant": [{"key": "c1", "label": "客户甲", "count": 3, "people": 1}],
                "members": [{"id": "c1", "name": "客户甲", "activities": [
                    {"date": "2026-10-01", "name": "读书会", "type": "沙龙活动", "subtype": "读书会",
                     "activity_type": "class", "course_key": "读书会", "role": "参与者", "teachers": "老师甲"}]}],
            },
        },
    }
    calls = []

    def analyze(request, data, *, export):
        assert export is True
        calls.append(data)
        return deepcopy(result)

    monkeypatch.setattr(api.principal_service, "analyze", analyze)
    monkeypatch.setattr(api, "audit", lambda *_a, **_kw: None)
    return result, calls


@pytest.mark.parametrize("group,view,expected", [
    ("traffic", "traffic", [("总数", "引流人数", "40 人"), ("引流人", "老师甲", "40 人"),
                            ("会员身份", "年卡", "5 人"), ("升单情况", "已升单", "2 人"),
                            ("跟进阶段", "未配置", "0 人"), ("流量来源", "自然到店", "3 人"),
                            ("客户标签", "标签甲", "2 人")]),
    ("invite", "invite_arrivals", [("总数", "邀约到店", "3 次"), ("总数", "发起邀约", "4 次"),
                                    ("邀约人 · 发起邀约", "接待甲", "4 次"), ("引流人 · 邀约到店", "老师甲", "3 次")]),
    ("invite", "invite_initiated", [("到店状态", "取消", "1 次"), ("到店状态", "未到店", "1 次"),
                                     ("到店状态", "已到店", "1 次")]),
    ("courses", "", [("总数", "课程数", "6 场"), ("总数", "课时数", "12 课时"),
                       ("活动类型", "沙龙活动", "6 场 · 12 课时"), ("沙龙类型 · 沙龙活动", "俱乐部 · 读书会", "6 场 · 12 课时"),
                       ("课程老师", "老师甲", "6 场 · 12 课时"), ("课程老师", "老师乙", "6 场 · 12 课时")]),
    ("deals", "", [("总数", "成交量", "9 笔"), ("总数", "总成交人数", "5 人"),
                     ("付费项目", "会员卡", "6 笔"), ("付费项目", "其他项目", "3 笔"),
                     ("卡种 · 会员卡", "年卡", "4 笔"), ("卡种 · 会员卡", "体验卡", "2 笔"),
                     ("购买类型", "首购", "5 笔"), ("购买类型", "复购", "4 笔")]),
    ("participation", "", [("总数", "活动参与", "3 次 · 1 人；案主 1 次 / 1 人；参与者 2 次 / 1 人")]),
])
def test_export_has_all_dropdown_counts_above_details(client, export_data, group, view, expected):
    _, calls = export_data
    response = client.post("/api/principal/export", json={
        "export_group": group, "export_view": view, "page_size": 1,
    })
    assert response.status_code == 200, response.text
    sheet = load_workbook(BytesIO(response.content)).active
    assert [cell.value for cell in sheet[1]][:3] == ["统计分类", "统计项目", "数量"]
    rows = list(sheet.values)
    details_index = next(index for index, row in enumerate(rows)
                         if row[0] in {"课程日期", "到店日期", "昵称", "发起邀约"} or row[:2] == ("参与者", "日期"))
    for category, label, value in expected:
        if category == "总数":
            assert (category, label, value) in [row[:3] for row in rows[:details_index]]
            continue
        matches = [(index, row) for index, row in enumerate(rows[:details_index]) if row[0] == category]
        assert len(matches) == 1
        index, row = matches[0]
        assert f"{label} {value}" in row[1].split("\n")
        assert row[2] is None
        assert f"B{index + 1}:C{index + 1}" in {str(merged) for merged in sheet.merged_cells.ranges}
        assert sheet.cell(index + 1, 2).alignment.wrap_text is True
        assert sheet.row_dimensions[index + 1].height >= 24
    assert not any(row[0] == "本次导出" or row[1] == "明细记录" for row in rows[:details_index])
    assert rows[details_index - 1][0] is None
    assert len(rows[details_index + 1:]) == (2 if group in {"courses", "deals"} else 1)
    assert sheet.cell(details_index + 2, 1).alignment.wrap_text is True
    assert calls[0].include_overview is True
    assert calls[0].tab == ("orders" if group == "deals" else "overview")


def test_export_summary_uses_only_authorized_result_and_sanitizes_option_names(client, export_data):
    result, _ = export_data
    result["breakdown"]["deals"] = []
    result["breakdown"]["buys"] = []
    result["summary"].pop("交易笔数")
    result["summary"].pop("成交人数")
    response = client.post("/api/principal/export", json={"export_group": "deals"})
    rows = list(load_workbook(BytesIO(response.content)).active.values)
    assert ("总数", "成交量", "0 笔") in [row[:3] for row in rows]
    assert not any(row[0] in {"付费项目", "购买类型", "卡种 · 会员卡"} for row in rows)
    result["breakdown"]["traffic_filters"]["tag"] = [{"key": "=1+1", "label": "=1+1", "count": 2}]
    response = client.post("/api/principal/export", json={"export_view": "traffic"})
    sheet = load_workbook(BytesIO(response.content)).active
    cell = next(row[1] for row in sheet if row[0].value == "客户标签")
    assert cell.value == "'=1+1 2 人" and cell.data_type == "s"


def test_empty_export_retains_zero_totals_and_detail_headers(client, export_data):
    result, _ = export_data
    result.update(summary={}, breakdown={}, items=[], total=0)
    response = client.post("/api/principal/export", json={"export_group": "courses"})
    rows = list(load_workbook(BytesIO(response.content)).active.values)
    assert ("总数", "课程数", "0 场") in [row[:3] for row in rows]
    assert rows[-1][:2] == ("课程日期", "课程")


def test_non_card_export_keeps_original_first_row(client, export_data):
    response = client.post("/api/principal/export", json={"tab": "courses"})
    rows = list(load_workbook(BytesIO(response.content)).active.values)
    assert rows[0][:2] == ("课程日期", "课程")


def test_multiple_dropdown_options_share_one_cell_and_keep_all_counts(client, export_data):
    result, _ = export_data
    result["breakdown"]["deals"] = [
        {"key": f"product-{index}", "label": f"付费项目{index}", "count": index, "subtypes": []}
        for index in range(5)
    ]
    response = client.post("/api/principal/export", json={"export_group": "deals"})
    sheet = load_workbook(BytesIO(response.content)).active
    cells = [row[1] for row in sheet if row[0].value == "付费项目"]
    assert len(cells) == 1
    assert cells[0].value == "\n".join(f"付费项目{index} {index} 笔" for index in range(5))
    assert sheet.row_dimensions[cells[0].row].height >= 5 * 16
