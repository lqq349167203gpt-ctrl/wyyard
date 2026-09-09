from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from app.api import custom_analysis as custom_analysis_api
from app.models.custom_analysis import AnalysisComparisonGroup, AnalysisCondition, AnalysisPlan
from app.services import custom_analysis_service


def _row(customer_id: str, **overrides):
    row = {
        "id": customer_id,
        "nickname": f"客户{customer_id}",
        "name": "",
        "gender": "",
        "age": "",
        "member_type": "普通会员",
        "follow_up_status": "前期沟通中",
        "customer_tags": [],
        "traffic_source": "小红书",
        "referrer": "小王",
        "referrer_handler": "小李",
        "service_teacher": "老师A",
        "referral_date": "2026-08-01",
        "created_at": "2026-08-01",
        "invitation_dates": [],
        "invitation_created_dates": [],
        "inviter_names": [],
        "first_visit_date": "",
        "last_visit_date": "",
        "invitation_count": 0,
        "visit_count": 0,
        "activity_count": 0,
        "communication_count": 0,
        "total_consumption": 0,
        "created_in_period": True,
        "referred_in_period": True,
        "invitation_count_period": 0,
        "visit_count_period": 0,
        "arrival_count_period": 0,
        "activity_count_period": 0,
        "payment_count_period": 0,
        "payment_amount_period": 0,
        "payment_dates": [],
        "course_teachers": [],
        "visit_purpose": "",
        "trauma_history": "",
        "current_block": "",
        "work_info": "",
        "other_info": "",
        "_payment_amounts_by_project_period": {},
        "_payment_orders_by_project_period": {},
        "_payment_events_period": [],
        "_invitation_events_period": [],
        "_invitation_events_all": [],
        "_activity_events_period": [],
    }
    row.update(overrides)
    return row


def test_execute_plan_filters_and_builds_dynamic_cards(monkeypatch):
    rows = [
        _row("c1", total_consumption=800, customer_tags=["高意向"]),
        _row("c2", total_consumption=1200, customer_tags=["高意向", "晚间偏好"]),
        _row("c3", visit_count=2, total_consumption=300, traffic_source="抖音"),
    ]
    monkeypatch.setattr(custom_analysis_service, "build_customer_dataset", lambda *_args: rows)
    plan = AnalysisPlan(
        title="未到店低消费客户",
        total_card_title="目标客户",
        conditions=[
            AnalysisCondition(field="visit_count", operator="eq", value=0),
            AnalysisCondition(field="total_consumption", operator="lt", value=1000),
        ],
        card_dimension="customer_tags",
        columns=["nickname", "visit_count", "total_consumption"],
        sort_by="total_consumption",
        sort_order="desc",
    )

    result = custom_analysis_service.execute_plan(plan, "actor", page=1, page_size=20)

    assert result["total"] == 1
    assert [item["id"] for item in result["items"]] == ["c1"]
    assert result["cards"] == [
        {"key": "total_customers", "title": "目标客户", "count": 1, "unit": "人", "format": "number", "is_total": True},
        {"key": "dimension-0", "title": "高意向", "count": 1, "unit": "人", "format": "number", "is_total": False},
    ]


def test_execute_plan_keeps_empty_values_at_end_when_descending(monkeypatch):
    rows = [
        _row("c1", referral_date=""),
        _row("c2", referral_date="2026-08-02"),
        _row("c3", referral_date="2026-08-01"),
    ]
    monkeypatch.setattr(custom_analysis_service, "build_customer_dataset", lambda *_args: rows)
    plan = AnalysisPlan(sort_by="referral_date", sort_order="desc")

    result = custom_analysis_service.execute_plan(plan, "actor", page=1, page_size=20)

    assert [item["id"] for item in result["items"]] == ["c2", "c3", "c1"]


def test_execute_plan_filters_by_any_invitation_date(monkeypatch):
    rows = [
        _row("c1", invitation_dates=["2026-08-01", "2026-08-05"]),
        _row("c2", invitation_dates=["2026-08-10"]),
        _row("c3", invitation_dates=[]),
    ]
    monkeypatch.setattr(custom_analysis_service, "build_customer_dataset", lambda *_args: rows)
    plan = AnalysisPlan(
        conditions=[
            AnalysisCondition(
                field="invitation_dates",
                operator="between",
                value=["2026-08-03", "2026-08-06"],
            ),
        ],
        columns=["nickname", "invitation_dates"],
        sort_by="invitation_dates",
    )

    result = custom_analysis_service.execute_plan(plan, "actor", page=1, page_size=20)

    assert [item["id"] for item in result["items"]] == ["c1"]


def test_execute_plan_supports_any_logic_and_selected_metrics(monkeypatch):
    rows = [
        _row("c1", traffic_source="小红书", invitation_count_period=1, visit_count_period=1),
        _row("c2", traffic_source="抖音", payment_count_period=2, payment_amount_period=796),
        _row("c3", traffic_source="朋友推荐"),
    ]
    monkeypatch.setattr(custom_analysis_service, "build_customer_dataset", lambda *_args: rows)
    plan = AnalysisPlan(
        conditions=[
            AnalysisCondition(field="traffic_source", operator="eq", value="小红书"),
            AnalysisCondition(field="payment_count_period", operator="gt", value=0),
        ],
        condition_logic="any",
        metrics=["total_customers", "arrived_customers", "converted_customers", "payment_orders", "payment_amount"],
        card_dimension="none",
    )

    result = custom_analysis_service.execute_plan(plan, "actor", page=1, page_size=20)

    assert result["total"] == 2
    assert {card["key"]: card["count"] for card in result["cards"]} == {
        "total_customers": 2,
        "arrived_customers": 1,
        "converted_customers": 1,
        "payment_orders": 2,
        "payment_amount": 796,
    }


def test_arrival_visits_counts_repeat_arrivals_for_same_customer(monkeypatch):
    rows = [
        _row("c1", visit_count_period=2, arrival_count_period=3),
        _row("c2", visit_count_period=1, arrival_count_period=2),
        _row("c3"),
    ]
    monkeypatch.setattr(custom_analysis_service, "build_customer_dataset", lambda *_args: rows)
    plan = AnalysisPlan(
        metrics=["arrived_customers", "arrival_visits"],
        card_dimension="none",
    )

    result = custom_analysis_service.execute_plan(plan, "actor", page=1, page_size=20)

    assert {card["key"]: card["count"] for card in result["cards"]} == {
        "arrived_customers": 2,
        "arrival_visits": 5,
    }


def test_activity_participations_counts_repeat_activities_for_same_customer(monkeypatch):
    rows = [
        _row("c1", activity_count_period=3),
        _row("c2", activity_count_period=2),
        _row("c3"),
    ]
    monkeypatch.setattr(custom_analysis_service, "build_customer_dataset", lambda *_args: rows)
    plan = AnalysisPlan(
        metrics=["activity_customers", "activity_participations"],
        card_dimension="none",
    )

    result = custom_analysis_service.execute_plan(plan, "actor", page=1, page_size=20)

    assert {card["key"]: card["count"] for card in result["cards"]} == {
        "activity_customers": 2,
        "activity_participations": 5,
    }


@pytest.mark.parametrize(
    ("row_display_mode", "count_field", "first_count", "second_count", "expected_total"),
    [
        ("arrival_visits", "arrival_count_period", 3, 2, 5),
        ("activity_participations", "activity_count_period", 2, 1, 3),
    ],
)
def test_execute_plan_can_expand_customer_rows_by_occurrence(
    monkeypatch,
    row_display_mode,
    count_field,
    first_count,
    second_count,
    expected_total,
):
    rows = [
        _row("c1", **{count_field: first_count}),
        _row("c2", **{count_field: second_count}),
        _row("c3"),
    ]
    monkeypatch.setattr(custom_analysis_service, "build_customer_dataset", lambda *_args: rows)
    plan = AnalysisPlan(
        metrics=["total_customers"],
        card_dimension="none",
        row_display_mode=row_display_mode,
    )

    result = custom_analysis_service.execute_plan(plan, "actor", page=1, page_size=20)

    assert result["total"] == expected_total
    assert result["total_unit"] == "人次"
    assert [item["id"] for item in result["items"]].count("c1") == first_count
    assert [item["id"] for item in result["items"]].count("c2") == second_count
    assert len({item["_display_key"] for item in result["items"]}) == expected_total
    assert result["cards"][0]["count"] == 3


def test_arrival_rows_are_unique_per_customer_day_and_show_one_invitation_date(monkeypatch):
    rows = [
        _row(
            "c1",
            arrival_count_period=3,
            invitation_dates=["2026-08-03", "2026-08-06"],
            _arrival_events_display=[
                {
                    "visit_date": "2026-08-03",
                    "invitation_created_dates": ["2026-08-01"],
                    "inviter_names": ["小李"],
                    "arrived": True,
                    "cancelled": False,
                },
                {
                    "visit_date": "2026-08-03",
                    "invitation_created_dates": ["2026-08-02"],
                    "inviter_names": ["小王"],
                    "arrived": True,
                    "cancelled": False,
                },
                {
                    "visit_date": "2026-08-06",
                    "invitation_created_dates": ["2026-08-05"],
                    "inviter_names": ["小张"],
                    "arrived": True,
                    "cancelled": False,
                },
            ],
        ),
    ]
    monkeypatch.setattr(custom_analysis_service, "build_customer_dataset", lambda *_args: rows)
    plan = AnalysisPlan(
        metrics=["total_customers"],
        card_dimension="none",
        row_display_mode="arrival_visits",
        columns=["nickname", "invitation_dates"],
    )

    result = custom_analysis_service.execute_plan(plan, "actor", page=1, page_size=20)

    assert result["total"] == 2
    assert [item["invitation_dates"] for item in result["items"]] == [
        ["2026-08-03"],
        ["2026-08-06"],
    ]
    assert len({item["_display_key"] for item in result["items"]}) == 2


def test_split_comparison_uses_selected_metric(monkeypatch):
    rows = [
        _row("c1", traffic_source="小红书", payment_amount_period=398),
        _row("c2", traffic_source="小红书", payment_amount_period=5999),
        _row("c3", traffic_source="抖音", payment_amount_period=0),
    ]
    monkeypatch.setattr(custom_analysis_service, "build_customer_dataset", lambda *_args: rows)
    plan = AnalysisPlan(
        card_metric="payment_amount",
        card_dimension="traffic_source",
    )

    result = custom_analysis_service.execute_plan(plan, "actor", page=1, page_size=20)
    dimension_cards = [card for card in result["cards"] if card["key"].startswith("dimension-")]

    assert dimension_cards == [
        {"key": "dimension-0", "title": "小红书", "count": 6397.0, "unit": "元", "format": "currency", "is_total": False},
        {"key": "dimension-1", "title": "抖音", "count": 0.0, "unit": "元", "format": "currency", "is_total": False},
    ]


def test_payment_amount_split_by_project_reconciles_with_total(monkeypatch):
    rows = [
        _row(
            "c1",
            payment_amount_period=300,
            purchased_projects=["会员卡·体验会员", "内部课程·疗愈师课程"],
            _payment_amounts_by_project_period={"会员卡·体验会员": 100, "内部课程·疗愈师课程": 200},
        ),
        _row(
            "c2",
            payment_amount_period=50,
            purchased_projects=["会员卡·体验会员"],
            _payment_amounts_by_project_period={"会员卡·体验会员": 50},
        ),
    ]
    monkeypatch.setattr(custom_analysis_service, "build_customer_dataset", lambda *_args: rows)
    plan = AnalysisPlan(
        metrics=["payment_amount"],
        card_metric="payment_amount",
        card_dimension="purchased_projects",
    )

    result = custom_analysis_service.execute_plan(plan, "actor", page=1, page_size=20)
    metric_card = next(card for card in result["cards"] if card["key"] == "payment_amount")
    dimension_cards = [card for card in result["cards"] if card["key"].startswith("dimension-")]

    assert metric_card["count"] == 350
    assert {card["title"]: card["count"] for card in dimension_cards} == {
        "内部课程·疗愈师课程": 200.0,
        "会员卡·体验会员": 150.0,
    }
    assert sum(card["count"] for card in dimension_cards) == metric_card["count"]


@pytest.mark.parametrize("operator", ["eq", "contains"])
def test_payment_condition_scopes_amount_and_orders_to_matching_transactions(monkeypatch, operator):
    rows = [
        _row(
            "c1",
            payment_categories=["会员卡", "内部课程"],
            payment_count_period=2,
            payment_amount_period=300,
            _payment_events_period=[
                {
                    "purchased_projects": ["会员卡·体验会员"],
                    "payment_categories": ["会员卡"],
                    "payment_projects": ["体验会员"],
                    "payment_closers": ["小王"],
                    "payment_methods": ["微信"],
                    "payment_dates": ["2026-08-01"],
                    "amount": 100,
                    "project_label": "会员卡·体验会员",
                },
                {
                    "purchased_projects": ["内部课程·疗愈师课程"],
                    "payment_categories": ["内部课程"],
                    "payment_projects": ["疗愈师课程"],
                    "payment_closers": ["小李"],
                    "payment_methods": ["支付宝"],
                    "payment_dates": ["2026-08-02"],
                    "amount": 200,
                    "project_label": "内部课程·疗愈师课程",
                },
            ],
        ),
        _row(
            "c2",
            payment_categories=["会员卡"],
            payment_count_period=1,
            payment_amount_period=50,
            _payment_events_period=[
                {
                    "purchased_projects": ["会员卡·体验会员"],
                    "payment_categories": ["会员卡"],
                    "payment_projects": ["体验会员"],
                    "payment_closers": ["小王"],
                    "payment_methods": ["微信"],
                    "payment_dates": ["2026-08-03"],
                    "amount": 50,
                    "project_label": "会员卡·体验会员",
                },
            ],
        ),
    ]
    monkeypatch.setattr(custom_analysis_service, "build_customer_dataset", lambda *_args: rows)
    plan = AnalysisPlan(
        conditions=[AnalysisCondition(field="payment_categories", operator=operator, value="会员卡")],
        metrics=["total_customers", "converted_customers", "payment_orders", "payment_amount"],
        card_dimension="none",
    )

    result = custom_analysis_service.execute_plan(plan, "actor", page=1, page_size=20)

    assert result["total"] == 2
    assert {card["key"]: card["count"] for card in result["cards"]} == {
        "total_customers": 2,
        "converted_customers": 2,
        "payment_orders": 2,
        "payment_amount": 150,
    }
    assert {item["id"]: item["payment_amount_period"] for item in result["items"]} == {
        "c1": 100,
        "c2": 50,
    }


def test_course_teacher_condition_scopes_activity_metrics(monkeypatch):
    rows = [
        _row(
            "c1",
            course_teachers=["奥雅", "耀凯"],
            activity_count_period=2,
            _activity_events_period=[
                {"activity_types": ["内部课程"], "activity_names": ["课程A"], "course_teachers": ["奥雅"]},
                {"activity_types": ["沙龙活动"], "activity_names": ["活动B"], "course_teachers": ["耀凯"]},
            ],
        ),
    ]
    monkeypatch.setattr(custom_analysis_service, "build_customer_dataset", lambda *_args: rows)
    plan = AnalysisPlan(
        conditions=[AnalysisCondition(field="course_teachers", operator="eq", value="奥雅")],
        metrics=["total_customers", "activity_customers", "activity_participations"],
        card_dimension="none",
    )

    result = custom_analysis_service.execute_plan(plan, "actor", page=1, page_size=20)

    assert {card["key"]: card["count"] for card in result["cards"]} == {
        "total_customers": 1,
        "activity_customers": 1,
        "activity_participations": 1,
    }
    assert result["items"][0]["activity_count_period"] == 1
    assert result["items"][0]["activity_names"] == ["课程A"]


def test_inviter_condition_scopes_invitation_metrics(monkeypatch):
    rows = [
        _row(
            "c1",
            inviter_names=["奥雅", "耀凯"],
            invitation_count_period=2,
            visit_count_period=2,
            arrival_count_period=2,
            _invitation_events_period=[
                {"inviter_names": ["奥雅"], "visit_date": "2026-01-05", "arrived": True, "cancelled": False},
                {"inviter_names": ["耀凯"], "visit_date": "2026-01-08", "arrived": True, "cancelled": False},
            ],
        ),
    ]
    monkeypatch.setattr(custom_analysis_service, "build_customer_dataset", lambda *_args: rows)
    plan = AnalysisPlan(
        conditions=[AnalysisCondition(field="inviter_names", operator="eq", value="奥雅")],
        metrics=["invited_customers", "arrived_customers", "arrival_visits"],
        card_dimension="none",
    )

    result = custom_analysis_service.execute_plan(plan, "actor", page=1, page_size=20)

    assert {card["key"]: card["count"] for card in result["cards"]} == {
        "invited_customers": 1,
        "arrived_customers": 1,
        "arrival_visits": 1,
    }
    assert result["items"][0]["invitation_count_period"] == 1
    assert result["items"][0]["visit_count_period"] == 1
    assert result["items"][0]["arrival_count_period"] == 1


def test_invitation_created_date_filters_and_splits_by_inviter(monkeypatch):
    rows = [
        _row(
            "c1",
            invitation_created_dates=["2026-08-10", "2026-08-11"],
            _invitation_events_all=[
                {
                    "inviter_names": ["潘潘"],
                    "invitation_created_dates": ["2026-08-10"],
                    "visit_date": "2026-08-15",
                    "arrived": False,
                    "cancelled": False,
                },
                {
                    "inviter_names": ["娟娟"],
                    "invitation_created_dates": ["2026-08-11"],
                    "visit_date": "2026-08-20",
                    "arrived": False,
                    "cancelled": False,
                },
            ],
        ),
        _row(
            "c2",
            invitation_created_dates=["2026-08-10"],
            _invitation_events_all=[
                {
                    "inviter_names": ["潘潘"],
                    "invitation_created_dates": ["2026-08-10"],
                    "visit_date": "2026-09-01",
                    "arrived": False,
                    "cancelled": False,
                },
            ],
        ),
    ]
    monkeypatch.setattr(custom_analysis_service, "build_customer_dataset", lambda *_args: rows)
    plan = AnalysisPlan(
        conditions=[
            AnalysisCondition(
                field="invitation_created_dates",
                operator="between",
                value=["2026-08-10", "2026-08-10"],
            ),
        ],
        metrics=["invited_customers"],
        card_metric="invited_customers",
        card_dimension="inviter_names",
    )

    result = custom_analysis_service.execute_plan(plan, "actor", page=1, page_size=20)

    assert result["total"] == 2
    assert {card["title"]: card["count"] for card in result["cards"]} == {
        "邀约人数": 2,
        "潘潘": 2,
    }
    assert all(item["inviter_names"] == ["潘潘"] for item in result["items"])


def test_comparison_groups_use_independent_periods_and_conditions(monkeypatch):
    def dataset(_actor_id, date_from="", date_to="", _allowed_customer_ids=None):
        if date_from == "2026-01-01":
            return [
                _row("c1", referrer="奥雅", payment_count_period=1, payment_amount_period=100),
                _row("c2", referrer="耀凯", payment_count_period=1, payment_amount_period=900),
            ]
        return [
            _row("c3", referrer="奥雅", payment_count_period=1, payment_amount_period=800),
            _row("c4", referrer="耀凯", payment_count_period=2, payment_amount_period=300),
        ]

    monkeypatch.setattr(custom_analysis_service, "build_customer_dataset", dataset)
    plan = AnalysisPlan(
        analysis_mode="comparison",
        metrics=["total_customers", "payment_orders", "payment_amount"],
        comparison_groups=[
            AnalysisComparisonGroup(
                id="a",
                name="1月奥雅",
                date_from="2026-01-01",
                date_to="2026-01-31",
                conditions=[AnalysisCondition(field="referrer", operator="eq", value="奥雅")],
            ),
            AnalysisComparisonGroup(
                id="b",
                name="2月耀凯",
                date_from="2026-02-01",
                date_to="2026-02-28",
                conditions=[AnalysisCondition(field="referrer", operator="eq", value="耀凯")],
            ),
        ],
    )

    result = custom_analysis_service.execute_plan(plan, "actor", page=1, page_size=20)

    assert [group["name"] for group in result["comparison_groups"]] == ["1月奥雅", "2月耀凯"]
    assert [row["values"] for row in result["comparison_rows"]] == [[1, 1], [1, 2], [100.0, 300.0]]
    assert result["comparison_rows"][2]["difference"] == 200.0
    assert result["comparison_rows"][2]["difference_rate"] == 200.0


def test_payment_date_condition_matches_any_transaction_date():
    row = _row("c1", payment_dates=["2026-05-12", "2026-08-20"])

    assert custom_analysis_service._matches(
        row,
        AnalysisCondition(field="payment_dates", operator="between", value=["2026-08-01", "2026-08-31"]),
    )
    assert not custom_analysis_service._matches(
        row,
        AnalysisCondition(field="payment_dates", operator="between", value=["2026-06-01", "2026-06-30"]),
    )


def test_date_condition_can_inherit_plan_period():
    condition = AnalysisCondition(
        field="payment_dates",
        operator="eq",
        value=None,
        inherit_period=True,
    )
    row = _row("c1", payment_dates=["2026-05-12", "2026-08-20"])

    assert custom_analysis_service._matches(row, condition, "2026-08-01", "2026-08-31")
    assert not custom_analysis_service._matches(row, condition, "2026-06-01", "2026-06-30")
    assert custom_analysis_service._matches(row, condition, "", "")


def test_non_date_condition_cannot_inherit_plan_period():
    with pytest.raises(ValueError, match="仅日期条件可跟随统计周期"):
        AnalysisCondition(
            field="nickname",
            operator="eq",
            value=None,
            inherit_period=True,
        )


def test_local_parser_recognizes_common_conditions(monkeypatch):
    monkeypatch.setattr(custom_analysis_service.customer_service, "list_customers", lambda: [
        SimpleNamespace(
            id="c1",
            member_type="普通会员",
            traffic_source="小红书",
            referrer="小王",
            referrer_handler="小李",
            service_teacher="老师A",
        ),
    ])
    monkeypatch.setattr(
        custom_analysis_service.customer_tag_service,
        "visible_tags_by_customer",
        lambda _actor_id: {},
    )

    plan = custom_analysis_service._local_plan(
        "筛选小红书来源、没有到店、消费低于1万元的客户，按流量来源统计",
        "actor",
    )

    conditions = {(item.field, item.operator, item.value) for item in plan.conditions}
    assert ("traffic_source", "eq", "小红书") in conditions
    assert ("visit_count", "eq", 0) in conditions
    assert ("total_consumption", "lt", 10000.0) in conditions
    assert plan.card_dimension == "traffic_source"


def test_metadata_endpoint(client):
    response = client.get("/api/custom-analysis/metadata")
    assert response.status_code == 200
    metadata = response.json()
    assert any(item["value"] == "follow_up_status" for item in metadata["fields"])
    assert any(item["value"] == "payment_dates" and item["label"] == "成交日期" for item in metadata["fields"])
    invitation_date = next(item for item in metadata["fields"] if item["value"] == "invitation_dates")
    assert invitation_date["label"] == "邀约日期"
    assert invitation_date["group"] == "日期信息"
    assert invitation_date["value_type"] == "date"
    assert "between" in invitation_date["operators"]
    arrival_visits = next(item for item in metadata["metrics"] if item["value"] == "arrival_visits")
    assert arrival_visits["label"] == "实际到场人次"
    assert arrival_visits["unit"] == "人次"
    activity_participations = next(
        item for item in metadata["metrics"] if item["value"] == "activity_participations"
    )
    assert activity_participations["label"] == "参与活动人次"
    assert activity_participations["unit"] == "人次"
    invitation_created_date = next(
        item for item in metadata["fields"] if item["value"] == "invitation_created_dates"
    )
    assert invitation_created_date["label"] == "邀约创建日期"
    assert invitation_created_date["group"] == "日期信息"
    assert invitation_created_date["value_type"] == "date"
    assert any(
        item["value"] == "inviter_names" and item["label"] == "邀约人"
        for item in metadata["card_dimensions"]
    )
    assert not any(item["value"] == "created_customers" for item in metadata["metrics"])
    assert any(item["value"] == "referred_customers" and item["label"] == "新引流客户数" for item in metadata["metrics"])
    assert {
        "visit_purpose",
        "trauma_history",
        "current_block",
        "work_info",
        "other_info",
    }.issubset({item["value"] for item in metadata["column_fields"]})
    assert not any(item["value"] == "visit_purpose" for item in metadata["fields"])


def test_metadata_only_returns_permitted_sensitive_column_options(client, monkeypatch):
    permissions = {
        "scope": "all",
        "relations": {"referrer": True, "referrer_handler": True},
        "sensitive_fields": {
            "visit_purpose": True,
            "trauma_history": False,
            "current_block": False,
            "work_info": True,
            "other_info": False,
        },
        "detail_tabs": {"communication": True},
        "transaction_access": "detail",
    }
    monkeypatch.setattr(
        custom_analysis_api.customer_access_service,
        "get_customer_permissions",
        lambda _role: permissions,
    )

    response = client.get("/api/custom-analysis/metadata")

    assert response.status_code == 200
    sensitive_options = {
        item["value"]
        for item in response.json()["column_fields"]
        if item["value"] in custom_analysis_service.SENSITIVE_COLUMN_FIELDS
    }
    assert sensitive_options == {"visit_purpose", "work_info"}


def test_execute_rejects_sensitive_column_without_permission(client, monkeypatch):
    permissions = {
        "scope": "all",
        "relations": {"referrer": True, "referrer_handler": True},
        "sensitive_fields": {
            "visit_purpose": True,
            "trauma_history": False,
            "current_block": True,
            "work_info": True,
            "other_info": True,
        },
        "detail_tabs": {"communication": True},
        "transaction_access": "detail",
    }
    monkeypatch.setattr(
        custom_analysis_api.customer_access_service,
        "get_customer_permissions",
        lambda _role: permissions,
    )

    response = client.post("/api/custom-analysis/execute", json={
        "plan": {
            "title": "越权字段测试",
            "conditions": [],
            "card_dimension": "none",
            "columns": ["nickname", "trauma_history"],
            "sort_by": "nickname",
            "sort_order": "asc",
        },
        "page": 1,
        "page_size": 20,
    })

    assert response.status_code == 403
    assert response.json()["detail"] == "当前角色没有创伤经历查看权限，请移除该显示列"


def test_execute_endpoint_returns_matching_customer(client, sample_customer):
    created = client.post("/api/customers", json=sample_customer).json()
    try:
        response = client.post("/api/custom-analysis/execute", json={
            "plan": {
                "title": "指定客户",
                "total_card_title": "符合条件",
                "conditions": [
                    {"field": "nickname", "operator": "eq", "value": created["nickname"]},
                ],
                "card_dimension": "none",
                "columns": ["nickname", "member_type"],
                "sort_by": "nickname",
                "sort_order": "asc",
            },
            "page": 1,
            "page_size": 20,
        })
        assert response.status_code == 200
        assert response.json()["total"] == 1
        assert response.json()["items"][0]["id"] == created["id"]

        logs_response = client.get("/api/analysis-logs?record_type=analysis")
        assert logs_response.status_code == 200
        assert all(item["log_type"] == "analysis_executed" for item in logs_response.json()["items"])
        matching_logs = [
            item for item in logs_response.json()["items"]
            if item["config"].get("标题") == "指定客户"
        ]
        assert matching_logs
        assert matching_logs[0]["config"]["筛选条件"] == [
            {"字段": "昵称", "规则": "等于", "值": created["nickname"]},
        ]
        assert matching_logs[0]["config"]["结果人数"] == 1
    finally:
        client.delete(f"/api/customers/{created['id']}")


def test_export_endpoint_uses_current_columns_and_row_display_mode(client, sample_customer):
    created = client.post("/api/customers", json=sample_customer).json()
    try:
        response = client.post("/api/custom-analysis/export", json={
            "plan": {
                "title": "客户导出测试",
                "conditions": [
                    {"field": "nickname", "operator": "eq", "value": created["nickname"]},
                ],
                "card_dimension": "none",
                "columns": ["nickname", "member_type"],
                "sort_by": "nickname",
                "sort_order": "asc",
                "row_display_mode": "unique_customers",
            },
        })

        assert response.status_code == 200, response.text
        assert response.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        workbook = load_workbook(filename=BytesIO(response.content))
        worksheet = workbook["筛选结果"]
        assert [worksheet.cell(row=1, column=index).value for index in range(1, 3)] == ["昵称", "会员身份"]
        assert worksheet.cell(row=2, column=1).value == created["nickname"]

        logs_response = client.get("/api/analysis-logs?record_type=export")
        assert logs_response.status_code == 200
        matching_logs = [
            item for item in logs_response.json()["items"]
            if item["log_type"] == "analysis_exported"
            and item["config"].get("标题") == "客户导出测试"
        ]
        assert matching_logs
        assert matching_logs[0]["config"]["结果人数"] == 1
    finally:
        client.delete(f"/api/customers/{created['id']}")


def test_export_contains_selected_sensitive_customer_columns(monkeypatch):
    monkeypatch.setattr(
        custom_analysis_service,
        "build_customer_dataset",
        lambda *_args: [
            _row(
                "c1",
                nickname="小安",
                visit_purpose="改善睡眠",
                trauma_history="童年经历",
                current_block="关系压力",
                work_info="在职 · 产品经理",
                other_info="偏好晚间联系，沟通时请完整记录客户反馈，并保留后续跟进安排和需要重点关注的信息。",
            ),
        ],
    )
    plan = AnalysisPlan(
        columns=[
            "nickname",
            "visit_purpose",
            "trauma_history",
            "current_block",
            "work_info",
            "other_info",
        ],
        card_dimension="none",
        sort_by="nickname",
    )

    output, record_count = custom_analysis_service.build_analysis_export(plan, "actor")

    assert record_count == 1
    worksheet = load_workbook(filename=output)["筛选结果"]
    assert [worksheet.cell(row=1, column=index).value for index in range(1, 7)] == [
        "昵称",
        "到访目的",
        "创伤经历",
        "当下卡点",
        "工作情况",
        "其他信息",
    ]
    assert [worksheet.cell(row=2, column=index).value for index in range(1, 7)] == [
        "小安",
        "改善睡眠",
        "童年经历",
        "关系压力",
        "在职 · 产品经理",
        "偏好晚间联系，沟通时请完整记录客户反馈，并保留后续跟进安排和需要重点关注的信息。",
    ]
    assert all(worksheet.cell(row=2, column=index).alignment.wrap_text for index in range(1, 7))
    assert worksheet.row_dimensions[2].height > 22


def test_comparison_execute_log_keeps_each_group_conditions_and_result(client, sample_customer):
    created = client.post("/api/customers", json=sample_customer).json()
    try:
        response = client.post("/api/custom-analysis/execute", json={
            "plan": {
                "title": "客户方案对比日志",
                "analysis_mode": "comparison",
                "metrics": ["total_customers"],
                "card_metric": "total_customers",
                "card_dimension": "none",
                "columns": ["nickname"],
                "sort_by": "nickname",
                "sort_order": "asc",
                "comparison_groups": [
                    {
                        "id": "group-a",
                        "name": "符合昵称",
                        "date_from": "2026-08-01",
                        "date_to": "2026-08-31",
                        "condition_logic": "all",
                        "conditions": [
                            {"field": "nickname", "operator": "eq", "value": created["nickname"]},
                        ],
                    },
                    {
                        "id": "group-b",
                        "name": "不符合昵称",
                        "condition_logic": "any",
                        "conditions": [
                            {"field": "nickname", "operator": "eq", "value": "不存在的客户"},
                        ],
                    },
                ],
            },
            "page": 1,
            "page_size": 20,
        })
        assert response.status_code == 200, response.text

        logs_response = client.get("/api/analysis-logs?record_type=analysis")
        assert logs_response.status_code == 200
        log = next(
            item for item in logs_response.json()["items"]
            if item["config"].get("标题") == "客户方案对比日志"
        )
        assert "符合昵称 1人" in log["content"]
        assert "不符合昵称 0人" in log["content"]
        assert log["config"]["分析模式"] == "方案对比"
        assert log["config"]["各组人数合计"] == 1
        assert log["config"]["对比组"] == [
            {
                "名称": "符合昵称",
                "时间范围": "2026-08-01 至 2026-08-31",
                "条件关系": "全部符合",
                "筛选条件": [{"字段": "昵称", "规则": "等于", "值": created["nickname"]}],
                "结果人数": 1,
            },
            {
                "名称": "不符合昵称",
                "时间范围": "全部时间",
                "条件关系": "任意一条符合",
                "筛选条件": [{"字段": "昵称", "规则": "等于", "值": "不存在的客户"}],
                "结果人数": 0,
            },
        ]
    finally:
        client.delete(f"/api/customers/{created['id']}")


def test_analysis_template_crud(client):
    payload = {
        "name": "本月引流测试模板",
        "description": "查看本月引流、邀约和到店转化",
        "scope": "private",
        "plan": AnalysisPlan(card_dimension="none").model_dump(mode="json"),
    }
    created = client.post("/api/custom-analysis/templates", json=payload)
    assert created.status_code == 200
    assert created.json()["description"] == payload["description"]
    template_id = created.json()["id"]
    try:
        listed = client.get("/api/custom-analysis/templates")
        assert listed.status_code == 200
        assert any(item["id"] == template_id for item in listed.json())

        logs = client.get("/api/analysis-logs?record_type=template")
        assert logs.status_code == 200
        assert all(item["log_type"].startswith("template_") for item in logs.json()["items"])
        saved_log = next(
            item for item in logs.json()["items"]
            if item["log_type"] == "template_created"
            and item["config"].get("模板名称") == payload["name"]
        )
        assert saved_log["operator"] == "不闹"
        assert saved_log["config"]["模板简介"] == payload["description"]
        assert saved_log["config"]["可见范围"] == "仅自己可见"

        used = client.post(f"/api/custom-analysis/templates/{template_id}/use")
        assert used.status_code == 200
        assert used.json()["use_count"] == 1

        updated = client.patch(
            f"/api/custom-analysis/templates/{template_id}",
            json={"name": "本月引流漏斗", "description": "更新后的模板简介"},
        )
        assert updated.status_code == 200
        assert updated.json()["name"] == "本月引流漏斗"
        assert updated.json()["description"] == "更新后的模板简介"
    finally:
        deleted = client.delete(f"/api/custom-analysis/templates/{template_id}")
        assert deleted.status_code == 200
