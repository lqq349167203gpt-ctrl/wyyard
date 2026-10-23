from io import BytesIO
from types import SimpleNamespace

from openpyxl import load_workbook

from app.services import customer_export_service as export


def test_card_statistics_stay_on_main_sheet_and_separate_coarse_cards():
    items = [{"nickname": "甲", "purchase_summary": [
        {"type": "会员卡", "remaining": 4, "used": 2, "current_remaining": 999, "debt_count": 3},
        {"type": "会员卡", "remaining": 6, "used": 1, "current_remaining": 999, "debt_count": 3},
        {"type": "会员卡", "name": "粗门次卡", "remaining": 1, "used": 1},
        {"type": "觉醒游戏", "current_remaining": 5, "used": 2},
    ]}]
    wb = load_workbook(BytesIO(export.build_workbook(items, columns=[{"key": "purchase_summary", "label": "卡次统计"}])))
    assert wb.sheetnames == ["客户资料"]
    rows = list(wb.active.values)
    values = dict(zip(*rows))
    assert values["会员卡剩余次数"] == 10
    assert values["会员卡欠卡次数"] == 3
    assert values["粗门次卡剩余次数"] == 1
    assert values["觉醒游戏剩余次数"] == 5


def test_unlimited_and_expired_cards():
    items = [{"purchase_summary": [{"type": "会员卡", "remaining": None}, {"type": "会员卡", "remaining": 50, "expiry_date": "2000-01-01"}]}]
    wb = load_workbook(BytesIO(export.build_workbook(items, columns=[{"key": "purchase_summary", "label": "卡次统计"}])))
    assert list(wb.active.values)[1][0] == "无限"


def test_summary_only_no_detail_sheets_and_empty_details_not_created():
    item = {"activity_summary": [{"type": "沙龙活动", "date": "2026-01-01", "participated": True}, {"type": "沙龙活动", "withdrawn": True}], "follow_up": []}
    wb = load_workbook(BytesIO(export.build_workbook([item], columns=[{"key": "activity_summary", "label": "次数"}, {"key": "follow_up", "label": "跟进点"}])))
    assert wb.sheetnames == ["客户资料"]
    assert list(wb.active.values)[1][0] == 1


def test_export_preferences_are_account_scoped_and_recheck_permissions(monkeypatch):
    from app.api import customers
    from app.services import storage
    monkeypatch.setattr(storage, "load_item", lambda _, key: {"columns": ["nickname", "phone"], "referrers": ["甲"]} if key == "export:a" else None)
    monkeypatch.setattr(customers, "get_request_roles", lambda _: [])
    monkeypatch.setattr(export, "available_columns", lambda _: [{"key": "nickname"}])
    assert customers.customer_export_preferences(SimpleNamespace(state=SimpleNamespace(user_id="a")))["columns"] == ["nickname"]
    assert customers.customer_export_preferences(SimpleNamespace(state=SimpleNamespace(user_id="b"))) == {}
