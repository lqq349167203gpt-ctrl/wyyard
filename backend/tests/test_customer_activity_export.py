from types import SimpleNamespace as NS

import pytest
from app.api import customer_detail
from app.services import customer_activity_export_service as service
from openpyxl import load_workbook


def activity(key, day="2026-09-01", **kwargs):
    return {"activity_key": key, "activity_type": key.split(":")[0], "date": day, "start_time": "09:00",
            "type": "沙龙类型", "course_type": "读书会", "name": "探索", "role": "参与者", "host": "老师甲",
            "participated": True, **kwargs}


def test_export_uses_last_balance_per_card_and_counts_only_attended_courses():
    activities = [activity("class:last", "2026-09-03"), activity("class:first"),
                  activity("class:absent", participated=False), activity("class:withdrawn", withdrawn=True),
                  activity("gcs:owner", "2026-09-02", role="案主", type="觉醒游戏")]
    cards = {"a": NS(card_type="60次卡", total_count=60, remaining_count=57, effective_date="2026-09-01")}
    usage = [{"key": "class:first", "card_id": "a", "remaining_after": 59},
             {"key": "class:first#unit=2", "card_id": "a", "remaining_after": 58},
             {"key": "class:last", "card_id": "a", "remaining_after": 57}]
    rows = service.activity_rows(activities, usage, cards, {"class:first": 2, "class:last": 1, "gcs:owner": 1})
    assert len(rows) == len(activities)
    assert [row[7] for row in rows] == [3, 2, 1, 1, 0]
    assert rows[3][2:4] == ["读书会", 2]
    assert rows[3][6] == "60次卡：58次"
    assert rows[1][6] == "觉醒游戏：历史余额未记录"
    assert rows[0][6] == "60次卡：57次"
    assert rows[2][6] == "已退课（不计入累计）"
    assert rows[4][6] == "未记录到店（不计入累计）"


def test_missing_snapshots_are_not_replaced_by_current_balance_and_zero_is_kept():
    card = NS(card_type="体验会员", total_count=4, remaining_count=0, effective_date="2026-09-01")
    usage = [{"key": "class:a", "card_id": "a", "remaining_after": None},
             {"key": "class:b", "card_id": "a", "remaining_after": 0}]
    rows = service.activity_rows([activity("class:a"), activity("class:b"), activity("class:c")], usage, {"a": card}, {"class:c": 1})
    assert rows[2][6] == "体验会员：历史余额未记录"
    assert rows[1][6] == "体验会员：0次"
    assert rows[0][6] == "历史余额未记录"


def test_distinguishes_free_unlimited_coarse_and_debt():
    rows = service.activity_rows([activity("class:a", role="老师"), activity("class:b"), activity("class:c"), activity("class:d")], [
        {"key": "class:b", "benefit_type": "internal_course", "benefit_name": "疗愈师课程"},
        {"key": "class:c", "benefit_type": "coarse_door_offset", "benefit_name": "粗门次卡"},
        {"key": "class:d", "benefit_type": "membership_debt", "benefit_name": "预支扣卡", "remaining_after": -2},
    ], {}, {})
    assert [row[6] for row in rows] == ["预支扣卡：欠2次", "粗门次卡：按课抵扣（无余额）", "疗愈师课程：不限次", "不扣卡"]


def test_current_pool_keeps_all_types_but_does_not_sum_repeated_membership_pool():
    summary = [{"type": "会员卡", "name": name, "current_remaining": 7, "current_total": 64, "debt_count": 2} for name in ("体验会员", "60次卡")]
    summary += [{"type": "觉醒游戏", "current_remaining": 1, "current_total": 2},
                {"type": "情绪释放", "current_remaining": 0, "current_total": 1},
                {"type": "能量结", "current_remaining": 3, "current_total": 6},
                {"type": "其他项目", "name": "专项", "current_remaining": 4, "current_total": 5},
                {"type": "内部课程", "name": "疗愈师课程", "current_remaining": "不限", "current_total": "不限"}]
    pools = service.pool_rows(summary)
    assert sum(label == "会员卡" for label, _ in pools) == 1
    assert all(any(label == item["type"] for label, _ in pools) for item in summary)
    assert pools[0] == ["会员卡", "当前剩余7次 / 总64次 · 历史欠卡2次"]
    assert ["情绪释放", "当前剩余0次 / 总1次"] in pools
    assert ["内部课程", "当前剩余不限次 / 总不限次"] in pools


def test_pool_preserves_each_card_status_deductions_dates_and_debt_details():
    common = {"type": "会员卡", "current_remaining": 6, "current_total": 10, "grand_total": 50,
              "debt_count": 2, "debt_activities": [{"label": "读书会", "date": "2026-08-01", "count": 2}],
              "internal_course_deductions": 3, "total_purchased": 10, "remaining": 6,
              "effective_date": "2026-09-01", "expiry_date": "2026-10-31"}
    cards = [{**common, "name": "生效卡", "activity_deductions": 2, "manual_deductions": 2},
             {**common, "name": "未来卡", "effective_date": "2026-10-02"},
             {**common, "name": "过期卡", "expiry_date": "2026-09-01"},
             {**common, "name": "用完卡", "remaining": 0},
             {**common, "name": "退费卡", "voided": True, "remaining": "已退费"}]
    rows = service.pool_rows(cards, today="2026-10-01")
    assert len([label for label, _ in rows if label in {c['name'] for c in cards}]) == 5
    by_label = dict(rows)
    assert by_label["生效卡"] == "剩余6次 / 总10次 · 2026-09-01~2026-10-31 · 生效中"
    for name, status in (("未来卡", "未生效"), ("过期卡", "已过期"), ("用完卡", "无次数"), ("退费卡", "已退费")):
        assert by_label[name].endswith(status)
    assert by_label["会员卡扣卡说明"] == "总50次 · 生效卡扣卡2次 · 销卡2次 · 沙龙技能学习抵扣3次"
    assert by_label["读书会"] == "2026-08-01 · 欠卡2次"


def test_pool_preserves_all_other_purchase_details_without_merging_or_pagination():
    projects = [{"type": "其他项目", "name": "专项", "remaining": 3, "total_purchased": 5,
                 "current_remaining": 3, "current_total": 5, "activity_mode": "线上",
                 "effective_date": "2026-09-01", "expiry_date": "2026-12-01"} for _ in range(7)]
    projects += [{"type": "觉醒游戏", "remaining": 1, "current_remaining": 1, "current_total": 2,
                  "purchases": [{"purchase_count": 2, "remaining": 1, "effective_date": "2026-09-01", "expiry_date": "2026-11-01"},
                                {"purchase_count": 3, "remaining": 0, "expiry_date": "2026-08-01"}]},
                 {"type": "线下落地课程", "name": "落地课程", "attended_count": 4, "validity_value": 6}]
    rows = service.pool_rows(projects, today="2026-10-01")
    assert sum(label == "专项 · 线上" for label, _ in rows) == 7
    assert ["购买 2 次", "剩余1次 / 总2次 · 2026-09-01~2026-11-01 · 生效中"] in rows
    assert ["购买 3 次", "剩余0次 / 总3次 · -~2026-08-01 · 已过期"] in rows
    assert ["线下落地课程", "已上课4次"] in rows
    assert ["落地课程", "6个月 · -~不限 · 生效中"] in rows


def test_participation_summary_separates_actual_attendance_and_all_roles():
    activities = [activity("class:a"), activity("class:b", role="老师", participated=False),
                  activity("gcs:c", type="觉醒游戏", role="案主"),
                  activity("gcs:d", type="觉醒游戏", role="成就君", withdrawn=True),
                  activity("ics:e", type="内部课程", role="助教")]
    rows = service.participation_rows(activities)
    assert ["活动类型", "沙龙活动", 2, 1, 1, 0] in rows
    assert ["活动类型", "觉醒游戏", 2, 1, 0, 1] in rows
    assert ["参与身份", "老师", 1, 0, 1, 0] in rows
    assert ["参与身份", "案主", 1, 1, 0, 0] in rows
    assert ["参与身份", "成就君", 1, 0, 0, 1] in rows
    assert ["参与身份", "助教", 1, 1, 0, 0] in rows
    assert all(row[2] == sum(row[3:]) for row in rows)


def test_workbook_has_eight_columns_no_freeze_and_literal_text():
    content = service.workbook_bytes("=危险昵称", [["会员卡", "当前剩余7次 / 总64次"]], [
        ["2026-09-01", "=SUM(1,2)", "读书会", 1, "参与者", "老师甲", "体验会员：3次", 1],
    ])
    sheet = load_workbook(content).active
    header_row = next(row[0].row for row in sheet if row[0].value == "活动日期")
    assert [sheet.cell(header_row, col).value for col in range(1, 9)] == service.HEADERS
    assert sheet.cell(header_row + 1, 2).data_type == "s"
    assert sheet.freeze_panes is None
    assert sheet.max_column == 8


@pytest.mark.parametrize("activities_allowed,cards_allowed", [(False, True), (True, False), (False, False)])
def test_endpoint_requires_both_permissions(client, monkeypatch, activities_allowed, cards_allowed):
    monkeypatch.setattr(customer_detail, "get_customer_detail", lambda *_args, **_kwargs: {
        "customer": {"id": "c", "customer_access_permissions": {"detail_tabs": {
            "activities": activities_allowed, "card_statistics": cards_allowed}}}, "activities": [],
    })
    monkeypatch.setattr(customer_detail, "_build_purchase_summary", lambda _: pytest.fail("拒绝前不应读取卡池"))
    response = client.get("/api/customer-detail/c/activity-export")
    assert response.status_code == 403


def test_endpoint_exports_all_rows_and_records_audit(client, monkeypatch):
    logs = []
    monkeypatch.setattr(customer_detail, "get_customer_detail", lambda *_args, **_kwargs: {
        "customer": {"id": "c", "nickname": "测试客户", "customer_access_permissions": {"detail_tabs": {
            "activities": True, "card_statistics": True}}},
        "activities": [activity(f"class:{index:03}", role="老师", participated=index < 17) for index in range(150)],
    })
    monkeypatch.setattr(customer_detail, "_build_purchase_summary", lambda _: [])
    monkeypatch.setattr(customer_detail.membership_card_service, "list_activity_usage_records", lambda _: [])
    monkeypatch.setattr(customer_detail.membership_card_service, "list_debt_activity_usage_records", lambda _: [])
    monkeypatch.setattr(customer_detail.membership_card_service, "list_cards", lambda: [])
    monkeypatch.setattr(customer_detail.operation_log_service, "create_log", lambda data, **kw: logs.append((data, kw)))
    response = client.get("/api/customer-detail/c/activity-export")
    assert response.status_code == 200
    from io import BytesIO
    sheet = load_workbook(BytesIO(response.content)).active
    header_row = next(row[0].row for row in sheet if row[0].value == "活动日期")
    assert sheet.max_row - header_row == 150
    assert sheet.cell(header_row + 1, 8).value == 17
    assert sheet.cell(sheet.max_row, 8).value == 1
    assert any(row[0].value == "活动参与次数与身份次数汇总" for row in sheet)
    assert any([cell.value for cell in row[:6]] == ["参与身份", "老师", 150, 17, 133, 0] for row in sheet)
    assert len(logs) == 1 and "150场" in logs[0][0].content
    assert logs[0][1]["extra"]["method"] == "EXPORT"


def test_missing_customer_cannot_export(client):
    assert client.get("/api/customer-detail/nonexistent-export-customer/activity-export").status_code == 404
