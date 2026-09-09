import io
import uuid
from datetime import datetime, timezone
from urllib.parse import quote

from fastapi import APIRouter, Depends, Query, Request
from fastapi.responses import StreamingResponse

from app.middleware.jwt_auth import require_page_permission
from app.models.login_record import LoginRecord
from app.models.operation_log import OperationLogCreate
from app.services import (
    account_service,
    customer_access_service,
    customer_service,
    login_record_service,
    operation_log_service,
    service_teacher_customer_service,
)
from app.utils.request_context import get_client_ip, get_client_source

router = APIRouter(
    prefix="/api/service-teacher-customers",
    tags=["service-teacher-customers"],
    dependencies=[Depends(require_page_permission("service-teacher"))],
)


def _actor_name(request: Request) -> str:
    return (
        getattr(request.state, "user_owner", "")
        or getattr(request.state, "user_name", "")
        or ""
    ).strip()


def record_service_teacher_action(request: Request, content: str, export: bool = False):
    account = account_service.get_account(getattr(request.state, "user_id", ""))
    if not account:
        return
    source = get_client_source(request)
    if export:
        # 使用统计的业务操作直接读取操作日志，避免重复写入两条。
        operation_log_service.create_log(OperationLogCreate(section="服务老师", content=content), extra={
            "operator": account.owner or account.username,
            "operator_role": "、".join(account.roles or [account.role]),
            "source": source, "method": "EXPORT", "path": request.url.path,
            "ip": get_client_ip(request),
        })
    else:
        login_record_service._save(LoginRecord(
            id=str(uuid.uuid4()), event_type="page_view", account_id=account.id,
            username=account.username, owner=account.owner or "", role=account.role,
            source=source, ip=get_client_ip(request),
            page_path="/pages/service-teachers/index" if source == "miniprogram" else "/service-teachers",
            page_name="服务老师", content=content, created_at=datetime.now(timezone.utc),
        ))


@router.post("/export-audit")
def record_pc_export(data: dict, request: Request):
    request.state.skip_operation_log = True
    content = str(data.get("content", ""))[:1000]
    record_service_teacher_action(request, f"导出服务老师记录：{content}", export=True)
    return {"success": True}


def _xlsx_response(sheet_name: str, headers: list[str], rows: list[list], filename: str):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    worksheet.sheet_view.showGridLines = False
    border = Border(bottom=Side(style="thin", color="F0F0F0"))
    header_fill = PatternFill(start_color="F7F8FA", end_color="F7F8FA", fill_type="solid")
    for column, label in enumerate(headers, 1):
        cell = worksheet.cell(row=1, column=column, value=label)
        cell.font = Font(bold=True, color="4E535A")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
        worksheet.column_dimensions[get_column_letter(column)].width = 18
    for row_number, values in enumerate(rows, 2):
        for column, value in enumerate(values, 1):
            cell = worksheet.cell(row=row_number, column=column, value=value)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = border
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    disposition = f"attachment; filename=service_teacher.xlsx; filename*=UTF-8''{quote(filename)}"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": disposition},
    )


@router.get("/metadata")
def get_metadata(request: Request):
    all_customers = customer_service.list_customers()
    customers = customer_access_service.filter_customers(request, all_customers)
    current_teacher = _actor_name(request)
    teachers = service_teacher_customer_service.available_teachers(customers, current_teacher)
    return {
        "current_teacher": current_teacher,
        "teachers": teachers,
        "teacher_options": service_teacher_customer_service.teacher_options(teachers, all_customers),
    }


@router.get("/export-follow-ups")
def export_follow_ups(
    request: Request,
    service_teacher: str | None = Query(None),
    follow_up_filter: str = Query("inactive", pattern="^(inactive|active|all|inactive_30|active_30)$"),
    follow_up_definition: str = Query("follow_up", pattern="^(none|customer_info|follow_up|both)$"),
    follow_up_days: int = Query(30, ge=1, le=3650),
):
    customers = customer_access_service.filter_customers(request, customer_service.list_customers())
    teacher = (service_teacher or _actor_name(request)).strip()
    result = service_teacher_customer_service.list_teacher_customers(
        customers,
        teacher,
        follow_up_filter=follow_up_filter,
        follow_up_definition=follow_up_definition,
        follow_up_days=follow_up_days,
        page=1,
        page_size=max(1, len(customers)),
    )
    rows = [
        [
            teacher,
            item["nickname"] or item["name"] or "-",
            item["member_type"] or "-",
            item["follow_up_status"] or "-",
            item["latest_customer_info_content"] or "-",
            item["latest_customer_info_by"] or "-",
            item["latest_customer_info_at"] or "-",
            item["latest_follow_up_content"] or "-",
            item["latest_follow_up_by"] or "-",
            item["latest_follow_up_at"] or "-",
            f"近{follow_up_days}天已录入" if item["is_active"] else f"近{follow_up_days}天未录入",
        ]
        for item in result["items"]
    ]
    definition_label = {
        "none": "客户信息或跟进点最近一项",
        "both": "客户信息和跟进点",
        "customer_info": "客户信息",
        "follow_up": "跟进点",
    }[follow_up_definition]
    normalized_filter = {"inactive_30": "inactive", "active_30": "active"}.get(follow_up_filter, follow_up_filter)
    filter_label = {"inactive": f"近{follow_up_days}天未录入", "active": f"近{follow_up_days}天已录入", "all": "全部客户"}[normalized_filter]
    record_service_teacher_action(request, f"导出跟进记录：服务老师 {teacher}；{definition_label}；{filter_label}；共{len(rows)}人", export=True)
    return _xlsx_response(
        "跟进记录",
        ["服务老师", "客户昵称", "会员身份", "跟进阶段", "最近客户信息", "客户信息录入人", "客户信息录入时间", "最近跟进点", "跟进点录入人", "跟进点录入时间", "当前状态"],
        rows,
        f"服务老师跟进记录_{teacher or '未选择'}.xlsx",
    )


@router.get("/export-courses")
def export_courses(
    request: Request,
    service_teacher: str | None = Query(None),
    teacher_id: str = Query(""),
    date_from: str | None = Query(None),
    date_to: str | None = Query(None),
    all_dates: bool = Query(False),
    activity_type: str = Query("all"),
):
    from app.api.statistics import get_course_statistics

    teacher = (service_teacher or _actor_name(request)).strip()
    result = get_course_statistics(
        date_from=date_from,
        date_to=date_to,
        all_dates=all_dates,
        granularity="day",
        organization_id=None,
        activity_type=activity_type,
        course_subtype=None,
        teacher_id=teacher_id or None,
        request=request,
    )
    rows = []
    for course in result["courses"]:
        newcomers = "、".join(
            participant["nickname"]
            for participant in course["participants"]
            if participant["identity_group"] == "新人"
        ) or "-"
        existing = "、".join(
            participant["nickname"]
            for participant in course["participants"]
            if participant["identity_group"] != "新人"
        ) or "-"
        course_time = course["start_time"]
        if course["end_time"]:
            course_time = f"{course_time}~{course['end_time']}"
        rows.append([
            course["date"] or "-",
            course_time or "-",
            course["name"] or "-",
            course["activity_type_label"] or "-",
            course["class_hours"],
            "、".join(course["teachers"]) or "-",
            course["participant_count"],
            newcomers,
            existing,
        ])
    record_service_teacher_action(request, f"导出课程记录：服务老师 {teacher}；{result['date_from']} 至 {result['date_to']}；共{len(rows)}场", export=True)
    return _xlsx_response(
        "课程记录",
        ["上课日期", "上课时间", "课程", "课程类型", "课时", "老师/成就君", "参与人数", "新人名单", "老人名单"],
        rows,
        f"服务老师课程记录_{teacher or '未选择'}_{result['date_from']}_{result['date_to']}.xlsx",
    )


@router.get("")
def list_customers(
    request: Request,
    service_teacher: str | None = Query(None),
    follow_up_filter: str = Query("inactive", pattern="^(inactive|active|all|inactive_30|active_30)$"),
    follow_up_definition: str = Query("follow_up", pattern="^(none|customer_info|follow_up|both)$"),
    follow_up_days: int = Query(30, ge=1, le=3650),
    nickname: str = Query(""),
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
):
    customers = customer_access_service.filter_customers(request, customer_service.list_customers())
    teacher = (service_teacher or _actor_name(request)).strip()
    definitions = {"none": "客户信息或跟进点最近一项", "both": "客户信息和跟进点", "customer_info": "客户信息", "follow_up": "跟进点"}
    normalized_filter = {"inactive_30": "inactive", "active_30": "active"}.get(follow_up_filter, follow_up_filter)
    filters = {"inactive": f"近{follow_up_days}天未录入", "active": f"近{follow_up_days}天已录入", "all": "全部客户"}
    record_service_teacher_action(request, f"查询跟进记录：服务老师 {teacher}；包含 {definitions[follow_up_definition]}；{filters[normalized_filter]}；第{page}页")
    return service_teacher_customer_service.list_teacher_customers(
        customers,
        teacher,
        follow_up_filter=follow_up_filter,
        follow_up_definition=follow_up_definition,
        follow_up_days=follow_up_days,
        nickname=nickname,
        page=page,
        page_size=page_size,
    )
