"""客户活动导出：复用活动/权益流水，只展示可追溯余额，不反推历史快照。"""

from collections import defaultdict
from datetime import datetime
from functools import cmp_to_key
from io import BytesIO

HEADERS = ["活动日期", "活动名称", "活动类型", "课时数", "参与身份", "课程老师", "活动后剩余卡次", "累计已上课数"]


def pool_rows(summary: list[dict], today: str | None = None) -> list[list]:
    """完整呈现详情页卡次统计：汇总、逐笔卡片、状态及欠卡明细，不合并购买明细。"""
    today = today or datetime.now().strftime("%Y-%m-%d")
    rows = []

    def expired(item):
        return bool(item.get("expiry_date") and item["expiry_date"] < today)

    def pending(item):
        return bool(item.get("effective_date") and item["effective_date"] > today)

    def status(item, full=True):
        if item.get("voided"):
            return "已退费"
        if expired(item):
            return "已过期"
        if pending(item):
            return "未生效"
        if item.get("remaining") == 0:
            return "无次数"
        return "生效中" if full or item["type"] == "线下落地课程" else ""

    def compare(a, b):
        sa, sb = (2 if expired(a) else 1 if pending(a) else 0), (2 if expired(b) else 1 if pending(b) else 0)
        if sa != sb:
            return sa - sb
        field = "effective_date" if sa == 1 else "expiry_date"
        av, bv = a.get(field) or ("9999" if sa == 0 else ""), b.get(field) or ("9999" if sa == 0 else "")
        return ((av > bv) - (av < bv)) * (-1 if sa == 2 else 1)

    def overview(item):
        if item["type"] == "线下落地课程":
            return f"已上课{item.get('attended_count', 0)}次"
        remaining = item.get("current_remaining", item.get("effective_remaining", item.get("remaining", 0)))
        total = item.get("current_total", item.get("total_purchased", 0))
        debt = item.get("debt_count", 0)
        return f"当前剩余{remaining}次 / 总{total}次" + (f" · 历史欠卡{debt}次" if debt else "")

    def detail(label, item, balance, full=True):
        values = [balance, f"{item.get('effective_date') or '-'}~{item.get('expiry_date') or '不限'}", status(item, full)]
        rows.append([label or "—", " · ".join(value for value in values if value)])

    def debts(item):
        if not item.get("debt_count", 0):
            return
        rows.append(["历史欠卡", f"{item['debt_count']}次"])
        for debt in item.get("debt_activities", []):
            rows.append([debt.get("label") or "活动", f"{debt.get('date') or '-'} · 欠卡{debt['count']}次"])

    members = [item for item in summary if item["type"] == "会员卡"]
    if members:
        first = members[0]
        unlimited = any(not expired(item) and not pending(item) and item.get("remaining") == "不限" for item in members)
        rows.append(["会员卡", "当前剩余不限次 / 总不限次" if unlimited else overview(first)])
        total = first.get("grand_total", sum(item.get("total_purchased", 0) for item in members if isinstance(item.get("total_purchased"), (int, float))))
        notes = ["不限" if unlimited else f"总{total}次"]
        for item in members:
            count = item.get("unlimited_deductions", 0) if "不限" in (item.get("remaining"), item.get("total_purchased")) else item.get("activity_deductions", 0)
            if count and not item.get("voided"):
                notes.append(f"{item.get('name', '')}扣卡{count}次")
        manual = sum(item.get("manual_deductions", 0) for item in members)
        if manual:
            notes.append(f"销卡{manual}次")
        if first.get("internal_course_deductions"):
            notes.append(f"沙龙技能学习抵扣{first['internal_course_deductions']}次")
        rows.append(["会员卡扣卡说明", " · ".join(notes)])
        for item in sorted(members, key=cmp_to_key(compare)):
            remaining, total = item.get("remaining", 0), item.get("total_purchased", 0)
            balance = "已退费" if item.get("voided") else "不限" if "不限" in (remaining, total) else f"剩余{max(0, remaining)}次 / 总{total}次"
            detail(item.get("name"), item, balance)
        debts(first)

    others = [item for item in summary if item["type"] not in {"会员卡", "OH卡诊断"}]
    for item in sorted(others, key=lambda item: bool(expired(item) or item.get("remaining") == 0)):
        rows.append([item["type"], overview(item)])
        purchases = sorted(item.get("purchases", []), key=cmp_to_key(compare))
        if purchases:
            for purchase in purchases:
                label = purchase.get("name") if item["type"] in {"沙龙技能学习", "其他项目"} else f"购买 {purchase['purchase_count']} 次"
                if item["type"] == "其他项目":
                    label = f"{label or ''} · {item.get('activity_mode') or '线下'}"
                remaining = purchase.get("remaining")
                if not isinstance(remaining, (int, float)):
                    remaining = purchase["purchase_count"]
                balance = "不限" if item["type"] == "沙龙技能学习" else f"剩余{remaining}次 / 总{purchase['purchase_count']}次"
                detail(label, purchase, balance)
        else:
            label = item.get("name")
            if item["type"] == "其他项目":
                label = f"{label or ''} · {item.get('activity_mode') or '线下'}"
            if item["type"] == "沙龙技能学习" or item.get("remaining") == "不限":
                balance = "不限"
            elif item["type"] == "线下落地课程":
                balance = f"{item.get('validity_value') or 1}个月"
            else:
                balance = f"剩余{item.get('remaining', 0)}次 / 总{item.get('total_purchased', 0)}次"
            detail(label, item, balance, full=False)
        debts(item)
    return rows


def participation_rows(activities: list[dict]) -> list[list]:
    """同一份完整活动清单分别按活动类型和身份统计，实际参与口径与详情页一致。"""
    types = {key: [0, 0, 0, 0] for key in ("沙龙活动", "觉醒游戏", "情绪释放", "能量结", "沙龙技能学习")}
    roles = {key: [0, 0, 0, 0] for key in ("参与者", "案主", "老师", "成就君")}
    for item in activities:
        kind = "沙龙活动" if item.get("activity_type") == "class" else item.get("type") or "未分类"
        # 退课优先，保证实际参与、未记录到店和已退课之和等于记录总数。
        state = 3 if item.get("withdrawn") else 1 if item.get("participated") else 2
        for group, key in ((types, kind), (roles, item.get("role") or "未标明身份")):
            counts = group.setdefault(key, [0, 0, 0, 0])
            counts[0] += 1
            counts[state] += 1
    return [[dimension, key, *counts] for dimension, groups in (("活动类型", types), ("参与身份", roles)) for key, counts in groups.items()]


def activity_rows(activities: list[dict], usage: list[dict], cards: dict, hours: dict) -> list[list]:
    by_activity = defaultdict(list)
    for record in usage:
        by_activity[str(record.get("key") or "").split("#unit=", 1)[0]].append(record)
    rows, attended_count = [], 0
    for activity in sorted(activities, key=lambda row: (row.get("date", ""), row.get("start_time") or "", row.get("activity_key", ""))):
        # 导出范围与详情页一致；到店/退课只影响累计次数，不能过滤整条记录。
        attended = bool(activity.get("participated")) and not activity.get("withdrawn")
        if attended:
            attended_count += 1
        balances = {}
        for record in by_activity[activity["activity_key"]]:
            card_id = record.get("card_id")
            card = cards.get(card_id)
            kind = record.get("benefit_type")
            label = record.get("benefit_name") or (card.card_type if card else "会员卡")
            key = card_id or record.get("benefit_id") or label
            value = record.get("remaining_after")
            if kind == "coarse_door_offset":
                text = "按课抵扣（无余额）"
            elif kind == "internal_course" or (card and card.total_count is None and card.remaining_count is None):
                text = "不限次"
            elif value is None:
                text = "历史余额未记录"
            else:
                # 同一活动扣多次时取最后扣减后的余额，不把每个扣减单位的余额相加。
                previous = balances.get(key, {}).get("value")
                if isinstance(previous, (int, float)):
                    value = min(value, previous)
                text = f"欠{abs(value)}次" if value < 0 else f"{value}次"
            balances[key] = {"label": label, "text": text, "value": value}
        if activity.get("role") == "案主" and activity.get("activity_type") in {"gcs", "ers", "eks"}:
            label = {"gcs": "觉醒游戏", "ers": "情绪释放", "eks": "能量结"}[activity["activity_type"]]
            # 专项案主目前只有按活动计算的销卡，没有逐场余额快照，不能以今天的余额冒充。
            balances.setdefault(label, {"label": label, "text": "历史余额未记录"})
        labels = [item["label"] for item in balances.values()]
        balance_text = []
        for key, item in balances.items():
            label = item["label"]
            if labels.count(label) > 1 and key in cards:
                label += f"（{cards[key].effective_date}生效，卡号{key[:8]}）"
            balance_text.append(f"{label}：{item['text']}")
        untracked = (activity.get("role") == "参与者" and activity["activity_type"] in {"class", "gcs", "ers"}
                     and not activity.get("is_public_welfare") and hours.get(activity["activity_key"], 0) > 0)
        balance_display = "\n".join(balance_text) if balance_text else ("历史余额未记录" if untracked else "不扣卡")
        if activity.get("withdrawn"):
            balance_display = "已退课（不计入累计）"
        elif not attended:
            balance_display = "未记录到店（不计入累计）"
        rows.append([
            activity["date"], activity.get("name") or "—",
            (activity.get("course_type") or "未分类") if activity["activity_type"] == "class" else activity["type"],
            hours.get(activity["activity_key"], 0), activity.get("role") or "—", activity.get("host") or "—",
            balance_display, attended_count,
        ])
    # 先正序计算历史累计，再倒序展示，不能用倒序行号代替累计已上课数。
    return rows[::-1]


def workbook_bytes(name: str, pools: list[list], rows: list[list], participation: list[list] | None = None) -> BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "客户活动记录"

    def cell(row, column, value):
        target = sheet.cell(row, column, value)
        if isinstance(value, str):
            target.data_type = "s"  # 名称、昵称等一律文本，避免 Excel 公式注入。
        target.font = Font(size=11, color="2B2F36")
        target.alignment = Alignment(vertical="center", wrap_text=True)
        return target

    cell(1, 1, f"{name} · 活动与卡池记录")
    sheet.merge_cells("A1:H1")
    sheet.row_dimensions[1].height = 28
    cell(2, 1, "卡次统计（与客户详情页一致，包含全部购买明细及历史欠卡）")
    sheet.merge_cells("A2:H2")
    for index, (label, value) in enumerate(pools or [["当前卡池", "暂无卡次记录"]], 3):
        cell(index, 1, label)
        sheet.merge_cells(start_row=index, start_column=1, end_row=index, end_column=2)
        cell(index, 3, value)
        sheet.merge_cells(start_row=index, start_column=3, end_row=index, end_column=8)
        sheet.row_dimensions[index].height = max(24, 18 * ((len(str(value)) + 79) // 80))
    note_row = 3 + max(1, len(pools))
    if participation:
        cell(note_row, 1, "活动参与次数与身份次数汇总")
        sheet.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=8)
        for column, label in enumerate(("统计维度", "类型 / 身份", "记录总数", "实际参与次数", "未记录到店", "已退课"), 1):
            cell(note_row + 1, column, label).fill = PatternFill("solid", fgColor="F7F8FA")
        for index, values in enumerate(participation, note_row + 2):
            for column, value in enumerate(values, 1):
                cell(index, column, value)
        note_row += len(participation) + 3
    cell(note_row, 1, "全部活动按日期、开始时间倒序展示；累计次数按发生顺序计算，一场计1次，未记录到店及已退课保留记录但不计累计。余额依据现有有效扣卡流水；无历史快照不推算，粗门按课抵扣不设余额。")
    sheet.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=8)
    sheet.row_dimensions[note_row].height = 36
    header_row = note_row + 2
    for column, label in enumerate(HEADERS, 1):
        target = cell(header_row, column, label)
        target.font = Font(bold=True, size=11, color="646A73")
        target.fill = PatternFill("solid", fgColor="F7F8FA")
    for index, values in enumerate(rows, header_row + 1):
        for column, value in enumerate(values, 1):
            cell(index, column, value).border = Border(bottom=Side(style="hair", color="ECEEF0"))
    for column, width in enumerate((16, 30, 24, 12, 14, 24, 48, 20), 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.auto_filter.ref = f"A{header_row}:H{max(header_row, sheet.max_row)}"
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
