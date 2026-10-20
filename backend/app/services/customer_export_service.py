"""客户资料导出：先筛选可见客户，再生成独立的 Excel 文件。"""

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models.customer import customer_referrers
from app.services import customer_access_service, customer_contact_service

# 字段由服务端统一给出，导出时再次校验，避免隐藏字段被手动请求带出。
COLUMNS = [
    ("nickname", "客户昵称", "基础资料"), ("name", "姓名", "基础资料"),
    ("gender", "性别", "基础资料"), ("age", "年龄", "基础资料"),
    ("member_type", "会员身份", "基础资料"), ("card_status", "卡次状态", "基础资料"),
    ("customer_tags", "客户标签", "基础资料"), ("service_teacher", "Sales", "基础资料"),
    ("customer_quality", "客户质量", "基础资料"), ("self_tags", "个人标签", "基础资料"),
    ("phone", "手机号", "联系来源"), ("wechat", "微信号", "联系来源"),
    ("referrer", "引流人", "联系来源"), ("referrer_handler", "承接人", "联系来源"),
    ("referral_date", "引流日期", "联系来源"), ("traffic_source", "流量来源", "联系来源"),
    ("traffic_post_name", "帖子名称", "联系来源"), ("traffic_source_detail", "来源详情", "联系来源"),
    ("tags", "到访目的", "客户档案"), ("basic_info", "创伤经历", "客户档案"),
    ("assessment", "当下卡点", "客户档案"), ("work_status", "工作状态", "客户档案"),
    ("work_description", "工作详情", "客户档案"), ("income_stability", "收入是否稳定", "客户档案"),
    ("personal_annual_income", "个人年收入", "客户档案"), ("household_annual_income", "家庭年收入", "客户档案"),
    ("other_info", "其他信息", "客户档案"),
    ("visit_count", "到店次数", "数据统计"), ("activity_count", "活动次数", "数据统计"),
    ("transaction_count", "交易笔数", "数据统计"),
    ("last_visit_date", "最近到访", "数据统计"), ("first_visit", "首次到访", "数据统计"), ("created_by", "创建人", "数据统计"),
    ("visit_need_tags", "来访需求标签", "跟进记录"), ("visit_need_detail", "来访需求详情", "跟进记录"),
    ("customer_info", "深度需求", "跟进记录"), ("follow_up", "跟进点", "跟进记录"),
    ("customer_feedback", "客户反馈", "跟进记录"), ("cancel_reason", "取消原因", "跟进记录"),
    ("communication_records", "沟通记录", "活动课程"), ("activities", "活动记录", "活动课程"),
    ("offline_course_records", "技能训练课程", "活动课程"),
    ("purchase_summary", "卡次统计", "卡次交易"), ("payment_records", "交易记录", "卡次交易"),
]
NOTE_COLUMNS = {"visit_need_tags", "visit_need_detail", "customer_info", "follow_up", "customer_feedback", "cancel_reason"}
RECORD_COLUMNS = NOTE_COLUMNS | {"communication_records", "activities", "offline_course_records", "purchase_summary", "payment_records"}
SECTION_COLUMNS = {"activities": "activities", "offline_course_records": "offline_course",
                   "purchase_summary": "purchase", "payment_records": "payment"}


def available_columns(roles) -> list[dict]:
    access = customer_access_service.get_customer_permissions(roles)
    denied = {field for permission, fields in customer_access_service.SENSITIVE_FIELD_MAP.items()
              if not access["sensitive_fields"][permission] for field in fields}
    if access["transaction_access"] == "none":
        denied.update({"transaction_count", "total_payment"})
    if access["transaction_access"] != "detail":
        denied.add("payment_records")
    if not customer_access_service.can_view_any_note(permissions=access):
        denied.add("first_visit")
    for column, permission in (("communication_records", "communication"), ("activities", "activities"),
                               ("offline_course_records", "offline_courses"), ("purchase_summary", "card_statistics")):
        if not access["detail_tabs"].get(permission):
            denied.add(column)
    for column in NOTE_COLUMNS:
        permission = "follow_up" if column == "cancel_reason" else column
        if not access["detail_tabs"].get(permission):
            denied.add(column)
    result = []
    for key, label, group in COLUMNS:
        if key in denied:
            continue
        if key in {"phone", "wechat"} and not customer_contact_service.can_access(roles, key, "copy"):
            label += "（脱敏）"
        result.append({"key": key, "label": label, "group": group,
                       "default": group in {"基础资料", "联系来源", "客户档案", "数据统计"}})
    return result


def prepare_items(customers, items, columns, roles, load_detail, communications=()):
    """仅查询所选的详情板块；沿用详情接口的客户范围与字段脱敏。"""
    selected = {column["key"] for column in columns}
    for customer, item in zip(customers, items):
        details = {}
        def detail_for(section):
            if section not in details:
                details[section] = load_detail(customer.id, section)
            return details[section]
        basic = customer.model_dump(mode="json") if hasattr(customer, "model_dump") else vars(customer)
        original = dict(item)
        item.update(basic)
        item.update(original)
        item["assessment"] = item.get("assessment") or item.get("core_situation") or ""
        for field in ("phone", "wechat"):
            if field in selected and not customer_contact_service.can_access(roles, field, "copy"):
                item[field] = customer_contact_service.mask_contact(field, item.get(field) or "")
        for key in selected.intersection({"total_payment", "first_visit"}):
            item[key] = detail_for("basic")["customer"][key]
        for key, section in SECTION_COLUMNS.items():
            if key in selected:
                item[key] = detail_for(section)[key]
                if key == "purchase_summary" and customer_access_service.get_customer_permissions(roles)["transaction_access"] == "none":
                    item[key] = _without_amounts(item[key])
        if "communication_records" in selected:
            item["communication_records"] = [record.model_dump(mode="json") for record in communications
                                             if (record.customer_id == customer.id if record.customer_id
                                                 else record.customer_nickname == customer.nickname)]
        if not selected.intersection(NOTE_COLUMNS):
            continue
        detail = detail_for("healing")
        for key in selected.intersection(NOTE_COLUMNS):
            item[key] = []
        for visit in detail["visit_records"]:
            for note in visit.get("visit_notes", []):
                category = note["category"]
                targets = {"visit_need_tags", "visit_need_detail"} if category == "visit_need" else {category}
                for key in targets.intersection(selected):
                    content = ("、".join(tag["name"] for tag in note.get("need_tags", [])) if key == "visit_need_tags"
                               else note.get("content", ""))
                    item[key].append({**note, "content": content, "visit_date": visit["visit_date"]})
            if "cancel_reason" in selected and visit.get("cancel_reason"):
                item["cancel_reason"].append({"visit_date": visit["visit_date"], "content": visit["cancel_reason"],
                                              "created_by": visit.get("created_by", "")})
        for note in detail.get("activity_participant_notes", []):
            category = note["category"]
            targets = {"visit_need_tags", "visit_need_detail"} if category == "visit_need" else {category}
            for key in targets.intersection(selected):
                content = ("、".join(tag["name"] for tag in note.get("need_tags", [])) if key == "visit_need_tags"
                           else note.get("content", ""))
                item[key].append({**note, "content": content})


def select_customers(customers, *, card_statuses=(), member_types=(), service_teachers=(),
                     referrers=(), referrer_handlers=(), date_from: date | None = None,
                     date_to: date | None = None):
    """同一项多选取任一匹配，不同项同时满足；引流日期包含起止当天。"""
    selections = (
        ("card_status", set(card_statuses)), ("member_type", set(member_types)),
        ("service_teacher", set(service_teachers)), ("referrer_handler", set(referrer_handlers)),
    )
    selected_referrers = set(referrers)
    matched = []
    for customer in customers:
        if any(values and (getattr(customer, field, "") or "").strip() not in values for field, values in selections):
            continue
        if selected_referrers and not selected_referrers.intersection(customer_referrers(customer)):
            continue
        if date_from or date_to:
            try:
                referral_date = date.fromisoformat(customer.referral_date)
            except (ValueError, TypeError):
                continue
            if date_from and referral_date < date_from or date_to and referral_date > date_to:
                continue
        matched.append(customer)
    return sorted(matched, key=lambda customer: str(customer.created_at or ""), reverse=True)


def build_workbook(items: list[dict], *, columns: list[dict]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "客户资料"
    sheet.append([column["label"] for column in columns])
    for item in items:
        row = []
        for column in columns:
            field = column["key"]
            value = item.get(field, "")
            if field == "customer_tags":
                value = "、".join(tag["name"] for tag in value or [])
            elif field == "referrer":
                value = "、".join(customer_referrers(item))
            elif field in RECORD_COLUMNS:
                value = f"共{len(value or [])}条，详见“{column['label']}”工作表"
            elif isinstance(value, list):
                value = "、".join(str(part) for part in value)
            elif isinstance(value, (date,)):
                value = value.isoformat()
            # 客户输入只能作为文本，不能在导出时执行公式。
            if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
                value = "'" + value
            row.append(value)
        sheet.append(row)
    record_fields = {
        "visit_date": "邀约日期", "date": "日期", "record_date": "记录日期", "deal_date": "成交日期",
        "start_time": "开始时间", "end_time": "结束时间", "type": "项目类型", "name": "名称",
        "course_name": "课程名称", "course_type": "课程类型", "activity_name": "活动名称",
        "activity_type": "活动类型", "activity_date": "活动日期", "role": "参与身份",
        "teacher": "老师", "host": "主持人", "title": "标题", "content": "内容", "result": "结果",
        "feedback_person": "反馈人", "created_by": "创建人", "creator": "记录人", "created_at": "录入时间",
        "updated_at": "更新时间", "effective_date": "生效日期", "expiry_date": "到期日期",
        "quantity": "数量", "total": "总次数", "total_count": "购买总次数", "used": "已用次数",
        "remaining": "剩余次数", "effective_remaining": "可用次数", "price": "金额", "amount": "金额",
        "closer_name": "成交人", "notes": "备注", "participated": "已参与", "withdrawn": "已退课",
        "voided": "已作废", "cancelled": "已取消", "cancellation_reason": "取消原因",
        "membership_deduction_count": "会员扣卡次数", "deduction_summary": "扣卡说明",
        "settlement_organization_name": "结算组织", "course_organization_name": "课程组织",
        "earliest_expiry": "最快到期日", "earliest_expiry_count": "最快到期次数",
        "current_remaining": "当前剩余", "current_total": "当前总次", "grand_total": "累计购买总次",
        "total_purchased": "累计购买次数", "total_amount": "总金额", "purchase_count": "购买次数",
        "manual_deductions": "手动销卡", "activity_deductions": "活动扣卡", "debt_count": "历史欠卡",
        "advance_deductions": "欠卡扣除", "internal_course_deductions": "技能学习抵扣",
        "unlimited_deductions": "不限次扣卡", "attended_count": "已上课数", "activity_mode": "活动模式",
        "validity_value": "有效期", "validity_unit": "有效期单位", "voided_at": "作废时间",
        "diagnosis_teacher": "诊断老师", "is_public_welfare": "公益活动",
    }

    def add_record_sheet(title, records):
        detail_sheet = workbook.create_sheet(title)
        fields = [key for key in record_fields if any(key in record for _, record in records)]
        detail_sheet.append(["客户昵称", *(record_fields[key] for key in fields)])
        for nickname, record in records:
            values = [nickname, *(record.get(key, "") for key in fields)]
            detail_sheet.append([_cell_value(value) for value in values])

    for column in columns:
        if column["key"] not in RECORD_COLUMNS:
            continue
        records = [(item.get("nickname") or item.get("name") or "", record)
                   for item in items for record in item.get(column["key"], [])]
        add_record_sheet(column["label"], records)
        if column["key"] == "purchase_summary":
            for key, title in (("purchases", "卡次购买明细"), ("debt_activities", "卡次欠卡明细")):
                children = [(nickname, {"type": record.get("type", ""), "name": record.get("name", ""), **child})
                            for nickname, record in records for child in record.get(key, [])]
                if children:
                    add_record_sheet(title, children)
    for worksheet in workbook:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cells in worksheet.iter_rows():
            for cell in cells:
                cell.font = Font(name="微软雅黑", size=11, bold=cell.row == 1,
                                 color="646A73" if cell.row == 1 else "2B2F36")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if cell.row == 1:
                    cell.fill = PatternFill("solid", fgColor="F2F3F5")
        for index, cell in enumerate(worksheet[1], 1):
            worksheet.column_dimensions[get_column_letter(index)].width = 42 if cell.value in {"内容", "备注", "工作详情", "创伤经历", "当下卡点", "其他信息"} else 20
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _cell_value(value):
    if isinstance(value, (dict, list)):
        value = "、".join(str(part) for part in value) if isinstance(value, list) else str(value)
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _without_amounts(value):
    if isinstance(value, list):
        return [_without_amounts(item) for item in value]
    if isinstance(value, dict):
        return {key: _without_amounts(item) for key, item in value.items() if key not in {"amount", "total_amount", "price"}}
    return value
